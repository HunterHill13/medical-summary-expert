#!/usr/bin/env python3
"""
tests/test_full_skill_diagnostics.py: Rigorous Stress-Testing & Bug Diagnosis Suite.
Tests:
1. Bug A: build_anki_apkg must auto-synthesize flashcards if 'anki_flashcards' is missing/empty.
2. Bug B: extract_clinical_numbers must capture Persian clinical units (درصد, میلی‌گرم, ساعت, روز, میلی‌متر جیوه)
          and Persian relational prefixes (کمتر از, بیش از, بالای, زیر, حداقل).
3. Bug C: add_comparison_table must never throw IndexError if row length exceeds headers length.
4. Bug D: build_docx_summary and build_markdown_summary must handle alternate pharmacology keys
          ('mechanism' vs 'mechanism_of_action', 'exam_pearl' vs 'pearl').
5. Bug E: null/empty diagnostics (best_initial_test is None or "") must not leave orphaned heading labels.
"""

import os
import sys
import json
import tempfile
import docx

SCRIPTS_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from build_summary_pack import (
    build_docx_summary,
    build_markdown_summary,
    build_anki_tsv,
    build_anki_apkg,
    add_comparison_table
)
from audit_summary_coverage import (
    extract_clinical_numbers,
    normalize_medical_text,
    audit_coverage
)

def test_bug_a_apkg_synthesis_when_empty():
    """Bug A: When anki_flashcards is omitted, .apkg must still be synthesized."""
    data = {
        "topic_title": "نارسایی قلبی تستی",
        "modules": [
            {
                "name": "HFrEF",
                "diagnostics": {
                    "best_initial_test": "نوار قلب و BNP",
                    "gold_standard": "اکوکاردیوگرافی داپلر"
                },
                "high_yield_pearls": ["کسر تخلیه‌ای زیر ۴۰ درصد است."]
            }
        ]
    }
    out_tsv = "dist/diag_test_a.tsv"
    out_apkg = "dist/diag_test_a.apkg"
    if os.path.exists(out_apkg):
        os.remove(out_apkg)
    if os.path.exists(out_tsv):
        os.remove(out_tsv)

    build_anki_tsv(data, out_tsv)
    build_anki_apkg(data, out_apkg)

    assert os.path.exists(out_tsv), "TSV was not created!"
    assert os.path.exists(out_apkg), "BUG A DETECTED: APKG was not created when anki_flashcards was omitted!"
    print("[PASS] Bug A Test Passed: APKG successfully synthesized when anki_flashcards is omitted.")

def test_bug_b_persian_units_and_prefixes():
    """Bug B: extract_clinical_numbers must extract numbers with Persian units and relational words."""
    persian_clinical_text = (
        "دوز اولیه ۵ میلی‌گرم خوراکی است. "
        "بیمار ظرف ۳ ساعت مجدداً پایش می‌شود. "
        "کسر تخلیه‌ای زیر ۴۰ درصد و فشار سیستولیک کمتر از ۹۰ میلی‌متر جیوه است."
    )
    norm = normalize_medical_text(persian_clinical_text)
    extracted = extract_clinical_numbers(norm)
    values_extracted = {item["value"] for item in extracted}

    print(f"[DEBUG] Extracted clinical items: {extracted}")
    # 5, 3, 40, 90 must all be extracted because they are associated with clinical units
    assert "5" in values_extracted, f"BUG B DETECTED: '5 میلی‌گرم' was dropped from clinical numbers! Extracted: {values_extracted}"
    assert "3" in values_extracted, f"BUG B DETECTED: '3 ساعت' was dropped from clinical numbers! Extracted: {values_extracted}"
    assert "40" in values_extracted, f"BUG B DETECTED: '40 درصد' was dropped! Extracted: {values_extracted}"
    assert "90" in values_extracted, f"BUG B DETECTED: '90 میلی‌متر جیوه' was dropped! Extracted: {values_extracted}"
    print("[PASS] Bug B Test Passed: Persian clinical units and prefixes correctly recognized.")

def test_bug_c_table_column_overflow():
    """Bug C: add_comparison_table must gracefully handle row data with more columns than headers without throwing IndexError."""
    doc = docx.Document()
    headers = ["ستون ۱", "ستون ۲"]
    # Row has 4 columns while headers has only 2
    rows = [["داده ۱", "داده ۲", "داده ۳ سرریز", "داده ۴ سرریز"]]
    try:
        add_comparison_table(doc, "جدول تستی سرریز", headers, rows)
        print("[PASS] Bug C Test Passed: Table overflow handled safely without IndexError.")
    except IndexError as e:
        raise AssertionError(f"BUG C DETECTED: add_comparison_table crashed on column overflow: {e}")

def test_bug_d_pharmacology_keys_interoperability():
    """Bug D: Ensure pharmacology_and_management works with both 'mechanism' and 'mechanism_of_action'."""
    data = {
        "topic_title": "تست فارماکولوژی",
        "modules": [
            {
                "name": "مبحث دارو",
                "pharmacology_and_management": [
                    {
                        "drug_class": "ACEi",
                        "mechanism": "مهار آنزیم مبدل آنژیوتانسین",
                        "exam_pearl": "سرفه خشک و آنژیوادم"
                    }
                ]
            }
        ]
    }
    doc_out = "dist/diag_test_d.docx"
    md_out = "dist/diag_test_d.md"
    build_docx_summary(data, doc_out)
    build_markdown_summary(data, md_out)

    with open(md_out, "r", encoding="utf-8") as f:
        md_content = f.read()

    assert "مهار آنزیم مبدل آنژیوتانسین" in md_content, "BUG D DETECTED: 'mechanism' key was omitted from Markdown output!"
    assert "سرفه خشک و آنژیوادم" in md_content, "BUG D DETECTED: 'exam_pearl' key was omitted from Markdown output!"
    print("[PASS] Bug D Test Passed: Alternate pharmacology keys rendered properly.")

def test_bug_e_null_diagnostics_no_orphan_headers():
    """Bug E: None/null diagnostic values must not create empty orphan labels."""
    data = {
        "topic_title": "تست تشخیص خالی",
        "modules": [
            {
                "name": "بیماری بدون تست قطعی",
                "diagnostics": {
                    "best_initial_test": "نوار قلب",
                    "gold_standard": None,
                    "key_labs": []
                }
            }
        ]
    }
    doc_out = "dist/diag_test_e.docx"
    build_docx_summary(data, doc_out)
    doc = docx.Document(doc_out)
    for p in doc.paragraphs:
        if "استاندارد طلایی و قطعی" in p.text:
            # If label exists, it must have text after it, not just the bare label!
            stripped = p.text.replace("⭐ استاندارد طلایی و قطعی (Gold Standard):", "").strip()
            assert len(stripped) > 0, "BUG E DETECTED: Orphaned 'Gold Standard' label with no content found!"
    print("[PASS] Bug E Test Passed: No orphan diagnostic labels generated.")

def test_bug_f_sdt_content_control_unrolling():
    """Bug F: Word documents with content controls (w:sdt) must not drop paragraphs or tables."""
    from parse_input_lecture import parse_docx
    from docx.oxml import OxmlElement

    doc = docx.Document()
    body = doc.element.body

    # Add regular paragraph
    doc.add_paragraph("پاراگراف عادی ۱")

    # Add SDT wrapped paragraph
    sdt = OxmlElement('w:sdt')
    content = OxmlElement('w:sdtContent')
    p_sdt = OxmlElement('w:p')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = "مطلب حیاتی درون کنترل ساختاریافته (SDT)"
    r.append(t)
    p_sdt.append(r)
    content.append(p_sdt)
    sdt.append(content)
    body.append(sdt)

    test_docx_path = "dist/diag_test_sdt.docx"
    doc.save(test_docx_path)

    parsed = parse_docx(test_docx_path, extract_images=False)
    assert "مطلب حیاتی درون کنترل ساختاریافته (SDT)" in parsed["full_text"], (
        "BUG F DETECTED: Word SDT content controls were skipped during document parsing!"
    )
    print("[PASS] Bug F Test Passed: SDT content control unrolling successfully extracts nested text.")

def test_bug_g_none_resilience_across_all_builders():
    """Bug G: All builders and audit functions must handle None values for all optional fields without crashing."""
    bad_data = {
        "topic_title": "مبحث با فیلدهای نال",
        "topic_title_en": None,
        "organ_system": None,
        "modules": [
            {
                "name": "مبحث تست نال",
                "name_en": None,
                "definition_and_etiology": None,
                "pathophysiology_cascade": None,
                "clinical_presentation": None,
                "diagnostics": None,
                "pharmacology_and_management": None,
                "differential_table": None,
                "high_yield_pearls": None,
                "instructor_exam_alerts": None,
                "contraindications_and_warnings": None,
                "scoring_criteria": None,
                "clinical_vignettes": None,
                "figures": None
            }
        ],
        "master_pharmacology_matrix": None,
        "last_minute_crash_sheet": None,
        "anki_flashcards": None
    }

    build_docx_summary(bad_data, "dist/diag_test_none.docx")
    build_markdown_summary(bad_data, "dist/diag_test_none.md")
    build_anki_tsv(bad_data, "dist/diag_test_none.tsv")
    build_anki_apkg(bad_data, "dist/diag_test_none.apkg")

    rep = audit_coverage("متن تست بالینی ۵ میلی‌گرم", bad_data)
    assert "verdict" in rep, "Audit failed on bad_data with None diagnostics!"
    print("[PASS] Bug G Test Passed: Zero crashes across DOCX, MD, TSV, APKG, and Audit with null fields.")

def test_bug_h_merge_batches_resilience():
    """Bug H: cmd_merge_batches must synthesize crash sheet safely even when fields are None."""
    from run_summary_pipeline import cmd_merge_batches
    import argparse

    test_batch_dir = "dist/test_batches"
    os.makedirs(test_batch_dir, exist_ok=True)

    manifest = {
        "file_name": "test_lecture.docx",
        "total_batches": 1,
        "batches": [
            {
                "batch_index": 1,
                "target_shard": "shard_01.json"
            }
        ]
    }
    with open(os.path.join(test_batch_dir, "manifest.json"), "w", encoding="utf-8") as fm:
        json.dump(manifest, fm)

    shard = {
        "topic_title": "تست مرج",
        "modules": [
            {
                "name": "سندروم تست",
                "definition_and_etiology": None,
                "pathophysiology_cascade": None,
                "clinical_presentation": None,
                "diagnostics": None,
                "pharmacology_and_management": None
            }
        ]
    }
    with open(os.path.join(test_batch_dir, "shard_01.json"), "w", encoding="utf-8") as fs:
        json.dump(shard, fs)

    out_merged = "dist/diag_test_merged.json"
    args = argparse.Namespace(batch_dir=test_batch_dir, out=out_merged)
    cmd_merge_batches(args)

    assert os.path.exists(out_merged), "Merged output was not created!"
    with open(out_merged, "r", encoding="utf-8") as f:
        merged_res = json.load(f)
    assert len(merged_res["last_minute_crash_sheet"]["rows"]) == 1, "Crash sheet row missing!"
    print("[PASS] Bug H Test Passed: merge-batches handled null fields gracefully.")

def test_bug_i_anki_cloze_model_type_and_card_generation():
    """Bug I: Verify Anki model_type is 1 (CLOZE) and every card produces at least one reviewable card."""
    import zipfile
    import sqlite3

    data = {
        "topic_title": "تست دک کلوز",
        "modules": [
            {
                "name": "مبحث پرل و کلوز",
                "high_yield_pearls": ["مروارید طلایی امتحانی بدون کلوز دستی"],
                "diagnostics": {
                    "best_initial_test": "نوار قلب",
                    "gold_standard": "اکوکاردیوگرافی"
                }
            }
        ],
        "anki_flashcards": [
            {
                "text": "کارت بدون براکت کلوز صریح",
                "extra": "توضیح تکمیلی",
                "tag": "Test"
            }
        ]
    }
    out_apkg = "dist/diag_test_cloze_type.apkg"
    build_anki_apkg(data, out_apkg)
    assert os.path.exists(out_apkg), "APKG was not created!"

    extract_dir = "dist/cloze_db_check"
    os.makedirs(extract_dir, exist_ok=True)
    with zipfile.ZipFile(out_apkg, 'r') as z:
        z.extract('collection.anki2', extract_dir)

    conn = sqlite3.connect(os.path.join(extract_dir, 'collection.anki2'))
    cur = conn.cursor()
    col = cur.execute('SELECT models FROM col').fetchone()[0]
    models = json.loads(col)
    for mid, mdata in models.items():
        assert mdata['type'] == 1, f"BUG I DETECTED: Anki model_type is {mdata['type']} instead of 1 (CLOZE)!"

    card_rows = cur.execute('SELECT id, nid FROM cards').fetchall()
    note_rows = cur.execute('SELECT id, flds FROM notes').fetchall()
    conn.close()

    assert len(card_rows) >= len(note_rows), f"BUG I DETECTED: Some notes generated zero cards! Notes: {len(note_rows)}, Cards: {len(card_rows)}"
    print(f"[PASS] Bug I Test Passed: Anki model_type is 1 (CLOZE) and all {len(note_rows)} notes generated {len(card_rows)} reviewable cards.")

def test_bug_j_partition_empty_and_markdown_tables():
    """Bug J: Verify partition on empty elements has word_count and markdown tables are parsed atomically."""
    from parse_input_lecture import partition_by_headings, parse_markdown_or_text
    import tempfile

    # 1. Empty elements partition
    empty_data = {"elements": [], "full_text": "پنج کلمه ساده برای تست"}
    batches = partition_by_headings(empty_data)
    assert len(batches) == 1, "Should have 1 batch"
    assert "word_count" in batches[0], "BUG J DETECTED: 'word_count' missing from empty elements batch!"
    assert batches[0]["word_count"] == 5, f"Expected 5 words, got {batches[0]['word_count']}"

    # 2. Markdown table preservation
    md_content = (
        "# عنوان مبحث\n\n"
        "متن مقدمه قبل از جدول.\n\n"
        "| بیماری | علامت |\n"
        "| --- | --- |\n"
        "| HFrEF | S3 Gallop |\n"
        "| HFpEF | S4 Gallop |\n\n"
        "متن بعد از جدول."
    )
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".md", delete=False) as tf:
        tf.write(md_content)
        tf_name = tf.name

    parsed = parse_markdown_or_text(tf_name)
    os.remove(tf_name)

    table_elements = [e for e in parsed["elements"] if e["type"] == "table"]
    assert len(table_elements) == 1, f"BUG J DETECTED: Markdown table was not parsed atomically! Found: {len(table_elements)}"
    assert "S3 Gallop" in table_elements[0]["markdown"], "Table markdown content incomplete!"
    print("[PASS] Bug J Test Passed: partition_by_headings and Markdown table grouping are robust.")

def test_bug_k_decimal_sentence_splitting_and_exact_number_matching():
    """Bug K: Verify decimal numbers are not broken by sentence splitting and numbers are matched with exact boundaries."""
    from audit_summary_coverage import (
        extract_clinical_numbers,
        extract_warnings_and_contraindications,
        normalize_medical_text,
        audit_coverage
    )

    # 1. Decimal extraction
    clinical_text = "دوز لیدوکائین 1.5 میلی‌گرم است. نسبت CTR > 0.5 و پتاسیم 3.5 mEq/L است."
    norm = normalize_medical_text(clinical_text)
    extracted = extract_clinical_numbers(norm)
    vals = {item["value"] for item in extracted}
    assert "1.5" in vals, f"BUG K DETECTED: '1.5' was broken into separate numbers! Extracted: {vals}"
    assert "0.5" in vals, f"BUG K DETECTED: '0.5' was broken into separate numbers! Extracted: {vals}"
    assert "3.5" in vals, f"BUG K DETECTED: '3.5' was broken into separate numbers! Extracted: {vals}"

    # 2. Warning with decimal
    warn_text = "در بیماران با کراتینین بالای 2.5 میلی‌گرم بر دسی‌لیتر مصرف متفورمین ممنوع است."
    warnings = extract_warnings_and_contraindications(normalize_medical_text(warn_text))
    assert len(warnings) == 1, "Warning with decimal was not captured!"
    assert "2.5" in warnings[0] and "کراتینین" in warnings[0], "Warning text was truncated at decimal point!"

    # 3. Exact boundary matching: 5 should NOT match 0.5 or 1.5
    src_text = "دوز اولیه ۵ میلی‌گرم است."
    sum_data = {
        "modules": [
            {
                "name": "تست عدم تطابق کاذب",
                "high_yield_pearls": ["نسبت CTR برابر 0.5 است ولی دوز ذکر نشده."]
            }
        ]
    }
    rep = audit_coverage(src_text, sum_data)
    # The source has 5 mg, but the summary only has 0.5. 5 must NOT match 0.5!
    assert rep["scores"]["numerical_coverage_pct"] == 0.0, (
        f"BUG K DETECTED: '5' falsely matched '0.5'! Coverage: {rep['scores']['numerical_coverage_pct']}"
    )
    print("[PASS] Bug K Test Passed: Decimals preserved and exact numeric boundaries enforced without false positive matches.")

def test_bug_l_empty_callout_box_suppression():
    """Bug L: add_callout_box and add_comparison_table must safely suppress empty tables."""
    from build_summary_pack import add_callout_box, add_comparison_table
    import docx

    doc = docx.Document()
    # Callout with empty list
    add_callout_box(doc, "کادر خالی", [], "85C1E9", "EBF5FB", "FFFFFF")
    # Comparison table with empty headers or rows
    add_comparison_table(doc, "جدول بدون هدر", [], [["داده"]])
    add_comparison_table(doc, "جدول بدون ردیف", ["هدر ۱"], [])

    assert len(doc.tables) == 0, f"BUG L DETECTED: Empty callout or comparison table created {len(doc.tables)} blank tables!"
    print("[PASS] Bug L Test Passed: Empty callouts and tables cleanly suppressed.")

def test_bug_m_word_break_and_tab_preservation():
    """Bug M: extract_text_including_txbx must preserve line breaks, tabs, and hyphens without fusing words."""
    from parse_input_lecture import extract_text_including_txbx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    xml = (
        f'<w:p {nsdecls("w")}>'
        '<w:r><w:t>دوز اول: ۵ میلی‌گرم</w:t></w:r>'
        '<w:r><w:br/></w:r>'
        '<w:r><w:t>دوز دوم: ۱۰ میلی‌گرم</w:t></w:r>'
        '<w:r><w:tab/></w:r>'
        '<w:r><w:t>حداکثر</w:t></w:r>'
        '<w:r><w:noBreakHyphen/></w:r>'
        '<w:r><w:t>دوز</w:t></w:r>'
        '</w:p>'
    )
    p_elem = parse_xml(xml)
    extracted = extract_text_including_txbx(p_elem)

    assert "میلی‌گرمدوز" not in extracted, f"BUG M DETECTED: Words fused across line break! Extracted: {repr(extracted)}"
    assert "دوز اول: ۵ میلی‌گرم" in extracted, "First line missing"
    assert "دوز دوم: ۱۰ میلی‌گرم" in extracted, "Second line missing"
    assert "حداکثر-دوز" in extracted, "Hyphenated word missing or broken"

    # Now verify clinical number extraction on this text
    norm = normalize_medical_text(extracted)
    nums = extract_clinical_numbers(norm)
    vals = {item["value"] for item in nums}
    assert "5" in vals and "10" in vals, f"BUG M DETECTED: Clinical numbers not extracted due to word fusion! Extracted: {vals}"
    print("[PASS] Bug M Test Passed: Word line breaks, tabs, and hyphens preserved without text fusion.")

def test_bug_n_master_matrices_word_markdown_parity():
    """Bug N: verify master_pharmacology_matrix and crash_sheet auto-synthesis parity and Mermaid label sanitation."""
    data = {
        "topic_title": "نارسایی قلبی و شوک",
        "topic_title_en": "Heart Failure & Shock",
        "organ_system": "قلب و عروق",
        "modules": [
            {
                "name": "HFrEF",
                "definition_and_etiology": ["کاهش قدرت انقباضی میوکارد"],
                "pathophysiology_cascade": ["کاهش برون‌ده قلبی [CO]", "فعال شدن RAAS [سیستم رنین]"],
                "clinical_presentation": [{"sign": "تنگی نفس ارتوپنه", "mechanism": "احتقان وریدی ریوی"}],
                "diagnostics": {"best_initial_test": "اکوکاردیوگرافی"},
                "pharmacology_and_management": [
                    {"drug_class": "ARNI", "mechanism_of_action": "مهار نپری‌لایزین", "indications": "HFrEF علامت‌دار"},
                    {"drug_class": "بتابلوکر", "mechanism_of_action": "مهار گیرنده بتا-1", "indications": "HFrEF پایدار"}
                ]
            }
        ]
    }
    out_md = "dist/diag_test_n.md"
    out_docx = "dist/diag_test_n.docx"

    build_markdown_summary(data, out_md)
    build_docx_summary(data, out_docx)

    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()

    # 1. Parity check: Both matrices synthesized
    assert "master_pharmacology_matrix" in data, "BUG N DETECTED: master_pharmacology_matrix was not auto-synthesized!"
    assert "last_minute_crash_sheet" in data, "BUG N DETECTED: last_minute_crash_sheet was not auto-synthesized!"
    assert "ماتریس جامع مقایسه داروشناسی بالینی" in md_text, "Master matrix heading missing from Markdown!"
    assert "کپسول طلایی شب امتحان" in md_text, "Crash sheet missing from Markdown!"

    # 2. Mermaid label sanitation: No square brackets inside quotes
    assert '["کاهش برون‌ده قلبی [CO]"]' not in md_text, "Mermaid square brackets were not sanitized to parentheses!"
    assert 'node0["کاهش برون‌ده قلبی (CO)"]' in md_text, f"Expected sanitized Mermaid node not found: {md_text}"

    # 3. Word check: Tables created
    doc = docx.Document(out_docx)
    assert len(doc.tables) >= 2, f"Word document missing synthesized matrix tables! Tables count: {len(doc.tables)}"

    print("[PASS] Bug N Test Passed: Master matrices auto-synthesized with 100% Word/Markdown parity and Mermaid sanitization.")

def test_bug_o_contraindication_false_negative_prevention():
    """Bug O: verify contraindication audit detects dropped warnings even if unrelated modules share generic words."""
    src = "در بیماران با نارسایی شدید کلیه، مصرف اسپیرونولاکتون اکیدا ممنوع است چون هایپرکالمی مرگبار دارد."

    # Summary A: has generic words "بیماران" and "کلیه" in an unrelated section, but dropped spironolactone warning
    sum_a = {
        "modules": [
            {
                "name": "بررسی کلیوی",
                "high_yield_pearls": ["بیماران با بیماری کلیه باید پایش شوند."]
            }
        ]
    }
    rep_a = audit_coverage(src, sum_a)
    assert len(rep_a["omitted_contraindications"]) == 1, (
        f"BUG O DETECTED: Dropped contraindication was not detected due to generic word match! "
        f"Omissions: {rep_a['omitted_contraindications']}"
    )
    assert rep_a["verdict"] == "WARNING_OMISSIONS_DETECTED", "Verdict should fail on omitted contraindication!"

    # Summary B: correctly preserved the spironolactone contraindication
    sum_b = {
        "modules": [
            {
                "name": "داروشناسی قلب",
                "contraindications_and_warnings": [
                    "مصرف اسپیرونولاکتون در نارسایی کلیه اکیدا ممنوع است."
                ]
            }
        ]
    }
    rep_b = audit_coverage(src, sum_b)
    assert len(rep_b["omitted_contraindications"]) == 0, (
        f"Preserved contraindication was falsely flagged as omitted! "
        f"Omissions: {rep_b['omitted_contraindications']}"
    )
    print("[PASS] Bug O Test Passed: Contraindication audit detects dropped warnings without false-negative bypass.")

def test_bug_p_thousands_and_arabic_decimal_commas():
    """Bug P: verify thousands separators (50,000 / ۱،۰۰۰،۰۰۰) and Arabic decimal commas (۱،۵) are normalized canonically without false omission alerts."""
    src = (
        "پلاکت بیمار کمتر از ۵۰،۰۰۰ سلول است. "
        "دوز لیدوکائین ۱،۵ میلی‌گرم است. "
        "ویروس لود بیمار بیش از ۱،۰۰۰،۰۰۰ کپی است."
    )
    sum_data = {
        "modules": [
            {
                "name": "تست هماتولوژی و ویروس",
                "high_yield_pearls": [
                    "پلاکت زیر 50000 سلول است.",
                    "دوز لیدوکائین 1.5 میلی‌گرم است.",
                    "ویروس لود بیش از 1000000 کپی است."
                ]
            }
        ]
    }
    report = audit_coverage(src, sum_data)
    assert report["scores"]["numerical_coverage_pct"] == 100.0, (
        f"BUG P DETECTED: Thousands or decimal commas caused false omission alerts! "
        f"Coverage: {report['scores']['numerical_coverage_pct']}%, Omitted: {report['omitted_clinical_numbers']}"
    )
    print("[PASS] Bug P Test Passed: Thousands and decimal commas normalized canonically.")

def test_bug_q_anki_contraindication_and_alert_synthesis():
    """Bug Q: verify auto-synthesis of Anki flashcards captures contraindications and instructor alerts."""
    from build_summary_pack import get_or_synthesize_flashcards, build_anki_tsv, build_anki_apkg

    data = {
        "topic_title": "فارماکولوژی پیشرفته قلب",
        "modules": [
            {
                "name": "متفورمین و وارفارین",
                "contraindications_and_warnings": [
                    "مصرف متفورمین در GFR زیر 30 اکیدا ممنوع است."
                ],
                "instructor_exam_alerts": [
                    "استاد تأکید کردند تداخل وارفارین و ریفامپین سؤال قطعی است."
                ]
            }
        ]
    }
    cards = get_or_synthesize_flashcards(data)
    tags = {c["tag"] for c in cards}

    assert any("Contraindications" in t for t in tags), f"BUG Q DETECTED: Contraindications not synthesized into Anki cards! Tags: {tags}"
    assert any("InstructorAlerts" in t for t in tags), f"BUG Q DETECTED: Instructor alerts not synthesized into Anki cards! Tags: {tags}"

    # Verify cloze syntax and extra field
    for c in cards:
        assert "{{c1::" in c["text"], f"Missing cloze syntax in card: {c['text']}"
        assert c["extra"], f"Missing extra field in card: {c}"

    out_tsv = "dist/diag_test_q.tsv"
    out_apkg = "dist/diag_test_q.apkg"
    build_anki_tsv(data, out_tsv)
    build_anki_apkg(data, out_apkg)
    assert os.path.exists(out_tsv) and os.path.exists(out_apkg), "TSV or APKG was not created for Bug Q!"

    print("[PASS] Bug Q Test Passed: Contraindications and instructor exam alerts synthesized into Anki flashcards.")

def test_bug_r_table_cell_textbox_preservation():
    """Bug R: table_to_markdown must capture text inside drawing shapes and textboxes within table cells."""
    from parse_input_lecture import table_to_markdown
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    # Row 0 (Headers)
    table.cell(0, 0).text = "مرحله بیماری"
    table.cell(0, 1).text = "اقدام درمانی"

    # Row 1, Cell 0: standard text
    table.cell(1, 0).text = "نارسایی قلبی کلاس IV"

    # Row 1, Cell 1: shape / drawing textbox inside paragraph
    p = table.cell(1, 1).paragraphs[0]
    p.text = ""  # empty run
    # Append drawing textbox XML to paragraph
    txbx_xml = parse_xml(r'''
        <w:r {} >
            <w:drawing>
                <wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
                    <a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                        <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                            <wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                                <wps:txbx>
                                    <w:txbxContent>
                                        <w:p><w:r><w:t>بستری فوری در CCU و اینوتروپ</w:t></w:r></w:p>
                                    </w:txbxContent>
                                </wps:txbx>
                            </wps:wsp>
                        </a:graphicData>
                    </a:graphic>
                </wp:inline>
            </w:drawing>
        </w:r>
    '''.format(nsdecls('w')))
    p._p.append(txbx_xml)

    # Verify standard cell.text misses the textbox
    assert table.cell(1, 1).text == "", "Pre-condition failed: python-docx cell.text should be empty for pure drawing textboxes"

    # Verify table_to_markdown captures the drawing textbox text
    md = table_to_markdown(table)
    assert "بستری فوری در CCU و اینوتروپ" in md, (
        f"BUG R DETECTED: Textbox inside table cell was omitted from markdown!\nMarkdown:\n{md}"
    )
    print("[PASS] Bug R Test Passed: Drawing textboxes inside table cells preserved in Markdown output.")

def test_bug_s_english_prose_stop_words_in_entity_recall():
    """Bug S: English sentence-starting prose words (When, Because, During, Other, First) must not trigger false omission warnings."""
    src = (
        "When managing Heart Failure, Especially during acute decompensation, First evaluate volume status. "
        "Because ARNI and empagliflozin improve survival, After stabilizing, initiate therapy. "
        "Other considerations include checking BNP and eGFR. Avoid verapamil."
    )
    # The summary is written in clinical Persian, containing only true medical acronyms/drugs (ARNI, empagliflozin, BNP, eGFR)
    sum_data = {
        "modules": [
            {
                "name": "مدیریت نارسایی قلبی",
                "pharmacology_and_management": [
                    {
                        "drug_class": "ARNI",
                        "mechanism": "مهار نپری‌لیسین و گیرنده آنژیوتانسین",
                        "exam_pearl": "درمان خط اول همراه با empagliflozin جهت کاهش مورتالیتی"
                    }
                ],
                "diagnostics": {
                    "best_initial_test": "سنجش سطح سرمی BNP و محاسبه eGFR",
                    "gold_standard": "اکوکاردیوگرافی داپلر"
                },
                "contraindications_and_warnings": [
                    "مصرف verapamil در این بیماران ممنوع است."
                ]
            }
        ]
    }
    rep = audit_coverage(src, sum_data)
    assert "When" not in rep["missing_entities"], f"BUG S DETECTED: 'When' treated as medical entity! Missing: {rep['missing_entities']}"
    assert "Because" not in rep["missing_entities"], f"BUG S DETECTED: 'Because' treated as medical entity! Missing: {rep['missing_entities']}"
    assert "During" not in rep["missing_entities"], f"BUG S DETECTED: 'During' treated as medical entity! Missing: {rep['missing_entities']}"
    assert "Other" not in rep["missing_entities"], f"BUG S DETECTED: 'Other' treated as medical entity! Missing: {rep['missing_entities']}"
    assert "First" not in rep["missing_entities"], f"BUG S DETECTED: 'First' treated as medical entity! Missing: {rep['missing_entities']}"
    assert "After" not in rep["missing_entities"], f"BUG S DETECTED: 'After' treated as medical entity! Missing: {rep['missing_entities']}"
    assert rep["scores"]["entity_recall_pct"] >= 85.0, (
        f"BUG S DETECTED: False low entity recall due to English prose! Score: {rep['scores']['entity_recall_pct']}%, Missing: {rep['missing_entities']}"
    )
    assert rep["verdict"] == "PASSED", f"Expected PASSED verdict, got: {rep['verdict']}"
    print("[PASS] Bug S Test Passed: English prose sentence starters correctly filtered, 100% genuine entity recall.")

def test_bug_t_html_and_pdf_generation_with_bidi_stamping():
    """Bug T: Direct HTML-to-PDF rendering with Playwright & Persian BiDi stamping."""
    from build_summary_pack import build_html_summary, build_pdf_summary
    
    sample_data = {
        "topic_title": "فیزیوپاتولوژی غدد درون‌ریز",
        "topic_title_en": "Endocrine Physiopathology",
        "organ_system": "فیزیوپاتولوژی غدد",
        "modules": [
            {
                "name": "سندرم کوشینگ و هیپرکورتیزولیسم",
                "definition_and_etiology": [
                    "افزایش مزمن سطح گلوکوکورتیکوئیدهای در گردش خون محیطی"
                ],
                "pathophysiology_cascade": [
                    "تومور هیپوفیز -> ترشح نابجای ACTH -> هایپرپلازی قشر آدرنال -> افزایش کورتیزول"
                ],
                "clinical_presentation": [
                    {"sign": "چاقی مرکزی، صورت ماه شکل (Moon face) و بوفالو هامپ", "mechanism": "لیپولیز محیطی و تجمع چربی احشایی"}
                ],
                "diagnostics": {
                    "best_initial_test": "تست مهار دگزامتازون 1 میلی‌گرم شبانه یا کورتیزول آزاد ادرار 24 ساعته",
                    "gold_standard": "نمونه‌برداری از سینوس پتروزال تحتانی (IPSS) جهت افتراق علل وابسته به ACTH"
                },
                "high_yield_pearls": [
                    "استریای بنفش عریض (> 1 cm) اختصاصی‌ترین علامت بالینی کوشینگ است."
                ]
            }
        ]
    }
    
    html_path = "dist/diag_test_t.html"
    pdf_path = "dist/diag_test_t.pdf"
    if os.path.exists(html_path):
        os.remove(html_path)
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
        
    build_html_summary(sample_data, html_path)
    assert os.path.exists(html_path), "HTML summary was not generated!"
    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()
    assert "@page" in html_content, "HTML missing @page CSS for PDF print!"
    assert "direction: rtl" in html_content, "HTML missing RTL direction!"
    assert "سندرم کوشینگ" in html_content, "Module title missing in HTML!"
    
    build_pdf_summary(html_path, pdf_path, course_name="فیزیوپاتولوژی غدد")
    assert os.path.exists(pdf_path), "PDF was not generated via Playwright!"
    assert os.path.getsize(pdf_path) > 1000, "Generated PDF is empty or corrupt!"
    print("[PASS] Bug T Test Passed: HTML-to-PDF generated via Playwright with Persian BiDi stamping.")

def test_bug_u_anti_pharmacology_skew_illness_script_integrity():
    """Bug U: Anti-Pharmacology-Skew Law & Illness Script Integrity Gate."""
    from audit_summary_coverage import audit_coverage
    
    # Source mentions subtle syndromes: شیهان and آپوپلکسی
    src = (
        "در هیپوپیتوئیتاریسم، سندرم شیهان ناشی از نکروز ایسکمیک پس از زایمان به دلیل خونریزی شدید است. "
        "همچنین آپوپلکسی هیپوفیز یک اورژانس حاد ناشی از خونریزی داخل آدنوم است که با سردرد ناگهانی شدید و افت بینایی تظاهر می‌کند. "
        "درمان دارویی جایگزینی لووتیروکسین 100 میکروگرم و هیدروکورتیزون است."
    )
    
    # Skewed summary that drops Sheehan and Apoplexy and only summarizes drug
    skewed_summary = {
        "modules": [
            {
                "name": "فارماکولوژی هورمون‌ها",
                "pharmacology_and_management": [
                    {
                        "drug_class": "لووتیروکسین",
                        "mechanism": "جایگزینی هورمون تیروئید T4 با نیمه عمر 7 روز و دوز 100 میکروگرم",
                        "exam_pearl": "مصرف صبح ناشتا با معده خالی"
                    }
                ]
            }
        ]
    }
    
    rep_skewed = audit_coverage(src, skewed_summary)
    assert not rep_skewed["scores"]["illness_script_compliant"], (
        "BUG U DETECTED: Skewed summary passed illness script integrity despite dropping Sheehan/Apoplexy!"
    )
    assert "شیهان" in rep_skewed["illness_script_audit"]["omitted_subtle_syndromes"]
    assert "آپوپلکسی" in rep_skewed["illness_script_audit"]["omitted_subtle_syndromes"]
    assert rep_skewed["verdict"] == "WARNING_OMISSIONS_DETECTED"
    
    # Sound summary preserving Sheehan and Apoplexy with 4 pillars
    sound_summary = {
        "modules": [
            {
                "name": "سندرم شیهان و آپوپلکسی هیپوفیز",
                "definition_and_etiology": [
                    "نکروز ایسکمیک غده هیپوفیز پس از زایمان به دنبال هموراژی شدید و افت فشار"
                ],
                "pathophysiology_cascade": [
                    "خونریزی حاد زایمان -> افت پرفیوژن هیپوفیز هیپرتروفیه -> انفارکتوس بافت سلار"
                ],
                "clinical_presentation": [
                    {"sign": "عدم توانایی شیردهی (Agalactorrhea) و آمنوره پس از زایمان", "mechanism": "نارسایی پرولاکتین و گنادوتروپین‌ها"}
                ],
                "diagnostics": {
                    "best_initial_test": "سنجش سطح سرمی هورمون‌های هیپوفیز (PRL, TSH, LH, FSH)",
                    "gold_standard": "MRI ناحیه هیپوفیز و زین ترکی"
                },
                "high_yield_pearls": [
                    "در آپوپلکسی هیپوفیز، تجویز فوری هیدروکورتیزون وریدی قبل از هر اقدام جراحی حیاتی است."
                ],
                "pharmacology_and_management": [
                    {
                        "drug_class": "لووتیروکسین و هیدروکورتیزون",
                        "mechanism": "جایگزینی هورمونی با دوز 100 میکروگرم",
                        "exam_pearl": "همیشه کورتون قبل از لووتیروکسین شروع شود تا بحران آدرنال رخ ندهد."
                    }
                ]
            }
        ]
    }
    
    rep_sound = audit_coverage(src, sound_summary)
    assert rep_sound["scores"]["illness_script_compliant"], (
        f"Sound summary failed illness script audit! Audit: {rep_sound['illness_script_audit']}"
    )
    assert len(rep_sound["illness_script_audit"]["omitted_subtle_syndromes"]) == 0
    print("[PASS] Bug U Test Passed: Anti-Pharmacology-Skew Law & Illness Script Integrity Gate enforces full disease script preservation.")

def test_bug_v_clinical_semiology_track_and_anki_synthesis():
    """Bug V: Track 4 Clinical Semiology & Physical Examination extraction and Anki card synthesis."""
    from build_summary_pack import (
        build_docx_summary,
        build_markdown_summary,
        build_html_summary,
        get_or_synthesize_flashcards
    )
    
    data = {
        "topic_title": "سمیولوژی غدد درون‌ریز",
        "organ_system": "غدد درون‌ریز",
        "clinical_semiology": [
            {
                "maneuver_or_sign": "مانور بلعیدن آب و لمس تیروئید",
                "technique": "قرارگیری پشت سر بیمار، لمس همزمان با بلع آب",
                "clinical_significance": "افتراق توده‌های تیروئید از سایر توده‌های قدام گردن",
                "pathology_association": "گواتر و ندول‌های تیروئید"
            },
            {
                "maneuver_or_sign": "علامت شوستک (Chvostek sign) و تروسو (Trousseau sign)",
                "technique": "ضربه به عصب فاسیال قدام گوش / باد کردن کاف فشارسنج بالاتر از سیستول به مدت 3 دقیقه",
                "clinical_significance": "اسپاسم کارپوپدال (Main d'accoucheur) ناشی از تحریک‌پذیری عصبی-عضلانی",
                "pathology_association": "هیپوکلسمی حاد و تتانی"
            }
        ],
        "modules": [
            {
                "name": "هیپوپاراتیروئیدیسم",
                "clinical_presentation": [{"sign": "تتانی و پارستزی", "mechanism": "کاهش کلسیم یونیزه"}],
                "diagnostics": {"best_initial_test": "سنجش کلسیم سرم", "gold_standard": "سنجش PTH"}
            }
        ]
    }
    
    # 1. Test Word generation
    out_docx = "dist/diag_test_v.docx"
    build_docx_summary(data, out_docx)
    assert os.path.exists(out_docx), "DOCX not generated!"
    
    # 2. Test Markdown generation
    out_md = "dist/diag_test_v.md"
    build_markdown_summary(data, out_md)
    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()
    assert "سمیولوژی و معاینه بالینی" in md_text, "Semiology heading missing in Markdown!"
    assert "علامت شوستک" in md_text, "Chvostek sign missing in Markdown!"
    
    # 3. Test HTML generation
    out_html = "dist/diag_test_v.html"
    build_html_summary(data, out_html)
    with open(out_html, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "semiology-card" in html_text or "semiology" in html_text, "Semiology styling missing in HTML!"
    assert "مانور بلعیدن آب" in html_text, "Maneuver missing in HTML!"
    
    # 4. Test Anki synthesis
    cards = get_or_synthesize_flashcards(data)
    semiology_cards = [c for c in cards if "Semiology" in c.get("tag", "")]
    assert len(semiology_cards) >= 2, f"Expected at least 2 semiology Anki cards, got: {len(semiology_cards)}"
    assert any("شوستک" in c["text"] for c in semiology_cards), "Chvostek sign missing from Anki cards!"
    
    print("[PASS] Bug V Test Passed: Clinical Semiology track rendered across DOCX, MD, HTML, and Anki decks.")

def test_bug_w_morning_exam_checklist_and_comparative_matrices():
    """Bug W: 30-Point Morning-of-Exam Golden Capsule and Comparative Matrices rendering."""
    from build_summary_pack import (
        build_docx_summary,
        build_markdown_summary,
        build_html_summary,
        ensure_master_matrices
    )
    
    data = {
        "topic_title": "اورژانس‌های غدد",
        "organ_system": "غدد درون‌ریز",
        "modules": [
            {
                "name": "طوفان تیروئیدی (Thyroid Storm)",
                "high_yield_pearls": [
                    "ترتیب احیا: ۱. پروپرانولول ۲. PTU یا متی‌مازول ۳. محلول ید لوگول (حداقل ۱ ساعت پس از PTU) ۴. هیدروکورتيزون"
                ],
                "contraindications_and_warnings": [
                    "تجویز آسپرین در طوفان تیروئیدی ممنوع است چون T4 متصل را آزاد می‌کند."
                ]
            }
        ],
        "comparative_matrices": [
            {
                "title": "جدول تطبیقی تیروئیدیت‌ها",
                "headers": ["نوع تیروئیدیت", "درد گردن", "ESR", "جذب ید (RAIU)", "درمان اصلی"],
                "rows": [
                    ["هاشیموتو (مزمن لنفوسیتی)", "ندارد (بدون درد)", "طبیعی", "متغیر / نرمال", "لووتیروکسین"],
                    ["دکورون (تحت حاد گرانولوماتوز)", "شدید و تیرکشنده به فک", "بسیار بالا (> 50)", "بسیار پایین (< 1%)", "NSAIDs / کورتون"]
                ]
            }
        ],
        "morning_exam_checklist": [
            "در طوفان تیروئیدی تجویز ید حتماً باید حداقل ۱ ساعت پس از PTU انجام شود.",
            "استریای بنفش عریض بالای ۱ سانتی‌متر اختصاصی‌ترین نشانه کوشینگ است."
        ]
    }
    
    # Check auto-synthesis if checklist omitted
    empty_checklist_data = {"topic_title": "تست", "modules": [{"name": "A", "high_yield_pearls": ["نکته ۱"]}]}
    ensure_master_matrices(empty_checklist_data)
    assert "morning_exam_checklist" in empty_checklist_data, "ensure_master_matrices failed to synthesize checklist!"
    assert len(empty_checklist_data["morning_exam_checklist"]) >= 1
    
    # Test Word
    out_docx = "dist/diag_test_w.docx"
    build_docx_summary(data, out_docx)
    assert os.path.exists(out_docx)
    
    # Test Markdown
    out_md = "dist/diag_test_w.md"
    build_markdown_summary(data, out_md)
    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()
    assert "چک‌لیست طلایی ۳۰ نکته‌ای" in md_text or "۳۰ نکته‌ای" in md_text
    assert "جدول تطبیقی تیروئیدیت‌ها" in md_text
    
    # Test HTML
    out_html = "dist/diag_test_w.html"
    build_html_summary(data, out_html)
    with open(out_html, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "checklist-box" in html_text or "morning_exam_checklist" in html_text or "کپسول" in html_text
    assert "جدول تطبیقی تیروئیدیت‌ها" in html_text
    
    print("[PASS] Bug W Test Passed: Morning exam checklist & comparative matrices rendered across all deliverables.")

def test_bug_x_pdf_designer_engine_dual_pass_and_geometry():
    """Bug X: Standalone MedicalPDFDesigner dual-pass rendering and print geometry."""
    from pdf_designer_engine import MedicalPDFDesigner
    import pypdf

    out_pdf = "dist/diag_test_x_designer.pdf"
    if os.path.exists(out_pdf):
        os.remove(out_pdf)

    designer = MedicalPDFDesigner(out_pdf, course_title="فیزیوپاتولوژی غدد • دانشکده پزشکی")
    sample_html = """<!DOCTYPE html>
    <html lang="fa" dir="rtl">
    <head><meta charset="utf-8">
    <style>
        @page { size: A4 portrait; margin-top: 24mm; margin-bottom: 20mm; margin-left: 17mm; margin-right: 17mm; }
        body { font-family: Tahoma, sans-serif; direction: rtl; text-align: right; }
    </style>
    </head>
    <body>
        <h1 style="color: #78281F;">بیماری‌های هیپوفیز و غدد درون‌ریز</h1>
        <p>آزمون تشخیصی و زنجیره پاتوفیزیولوژی آدنوم هیپوفیز و آپوپلکسی.</p>
    </body>
    </html>
    """
    designer.add_section("sec_hypo", sample_html, topic_title="آدنوم هیپوفیز و سندرم شیهان")
    res_path = designer.build(tmp_dir="dist/tmp_x_designer")

    assert os.path.exists(res_path), "PDF Designer failed to output file!"
    assert os.path.getsize(res_path) > 1000, "Output PDF is too small or corrupt!"

    reader = pypdf.PdfReader(res_path)
    assert len(reader.pages) == 1, "Expected exactly 1 page!"
    p0 = reader.pages[0]
    w = float(p0.mediabox.width)
    h = float(p0.mediabox.height)
    assert abs(w - 595.28) < 1.0, f"Page width mismatch: {w} != 595.28 pt"
    assert abs(h - 841.89) < 1.0, f"Page height mismatch: {h} != 841.89 pt"

    print("[PASS] Bug X Test Passed: MedicalPDFDesigner dual-pass hybrid rendering and A4 geometry verified.")

def test_bug_y_comprehensive_course_bundling():
    """Bug Y: Comprehensive Master Course Bundling (Cover + TOC + Chapters + 'دانشگاه بق' header)."""
    from pdf_designer_engine import MedicalPDFDesigner
    from run_summary_pipeline import cmd_bundle
    import argparse

    os.makedirs("dist/test_bundle_y", exist_ok=True)
    ch1_path = "dist/test_bundle_y/session_1.json"
    ch2_path = "dist/test_bundle_y/session_2.json"

    mock_ch1 = {
        "topic_title": "جلسه ۱: اختلالات محور هیپوتالاموس-هیپوفیز",
        "topic_title_en": "Hypothalamic-Pituitary Disorders",
        "organ_system": "غدد درون‌ریز",
        "institution": "دانشگاه بق",
        "modules": [
            {
                "name": "آدنوم هیپوفیز و پرولاکتینوما",
                "pathophysiology_chain": "افزایش ترشح هورمون رشد یا پرولاکتین -> تومور ماکروآدنوم -> فشار بر کیاسمای بینایی -> همی‌آنوپسی بای‌تمپورال.",
                "diagnostics": {
                    "best_initial_test": "سنجش سطح سرمی پرولاکتین ناشتا",
                    "gold_standard": "MRI هیپوفیز با تزریق گادولینیوم"
                },
                "treatment_protocol": {
                    "first_line": "آگونیست‌های دوپامین (کابرگولین یا بروموکریپتین)",
                    "definitive": "جراحی ترانس‌اسفنوئیدال در صورت شکست درمان دارویی"
                },
                "high_yield_pearls": ["خط اول درمان ماکروپرولاکتینوما برخلاف سایر آدنوم‌ها جراحی نیست، بلکه کابرگولین است."]
            }
        ],
        "anki_flashcards": [
            {
                "question": "خط اول درمان {{c1::ماکروپرولاکتینوما}} داروی {{c2::کابرگولین}} است.",
                "answer": "کابرگولین آگونیست دوپامین است."
            }
        ]
    }

    mock_ch2 = {
        "topic_title": "جلسه ۲: بیماری‌های غده تیروئید و گواتر",
        "topic_title_en": "Thyroid Disorders & Goiter",
        "organ_system": "غدد درون‌ریز",
        "institution": "دانشگاه بق",
        "modules": [
            {
                "name": "بیماری گریوز (Graves Disease)",
                "pathophysiology_chain": "آنتی‌بادی محرک گیرنده TSH (TSI/TRAb) -> تحریک مداوم سنتز هورمون‌های T3 و T4 -> هیپرتیروئیدی همراه گواتر منتشر.",
                "diagnostics": {
                    "best_initial_test": "سنجش سطح سرمی TSH و Free T4",
                    "gold_standard": "اسکن تیروئید با ید رادیواکتیو (RAIU) و تیتر TSI"
                },
                "treatment_protocol": {
                    "first_line": "متی‌مازول (در سه ماهه اول بارداری PTU)",
                    "definitive": "ید رادیواکتیو ۱۳۱ یا تیروئیدکتومی ساب‌توتال"
                },
                "high_yield_pearls": ["در تیروئیدوتوکسیکوز شدید، پروپرانولول علائم آدرنرژیک و تبدیل محیطی T4 به T3 را مهار می‌کند."]
            }
        ],
        "anki_flashcards": [
            {
                "question": "در سه ماهه اول بارداری برای درمان گریوز داروی {{c1::PTU}} و در ادامه {{c2::متی‌مازول}} ترجیح دارد.",
                "answer": "به علت خطر تراتوژنیسیتی متی‌مازول در ارگانوژنز."
            }
        ]
    }

    with open(ch1_path, "w", encoding="utf-8") as f:
        json.dump(mock_ch1, f, ensure_ascii=False, indent=2)

    with open(ch2_path, "w", encoding="utf-8") as f:
        json.dump(mock_ch2, f, ensure_ascii=False, indent=2)

    # Test cmd_bundle
    bundle_args = argparse.Namespace(
        inputs=[ch1_path, ch2_path],
        course_title="درسنامه جامع فیزیوپاتولوژی غدد",
        institution="دانشگاه بق",
        out_dir="dist/test_bundle_y/output",
        prefix="master_endocrine_booklet"
    )
    cmd_bundle(bundle_args)

    pdf_out = os.path.join(bundle_args.out_dir, "master_endocrine_booklet.pdf")
    docx_out = os.path.join(bundle_args.out_dir, "master_endocrine_booklet.docx")
    md_out = os.path.join(bundle_args.out_dir, "master_endocrine_booklet.md")
    apkg_out = os.path.join(bundle_args.out_dir, "master_endocrine_booklet_anki.apkg")
    tsv_out = os.path.join(bundle_args.out_dir, "master_endocrine_booklet_anki.tsv")

    assert os.path.exists(pdf_out), "Master PDF bundle not generated!"
    assert os.path.exists(docx_out), "Master DOCX bundle not generated!"
    assert os.path.exists(md_out), "Master Markdown bundle not generated!"
    assert os.path.exists(apkg_out), "Master APKG bundle not generated!"
    assert os.path.exists(tsv_out), "Master TSV bundle not generated!"

    import pypdf
    r = pypdf.PdfReader(pdf_out)
    assert len(r.pages) >= 3, f"Expected at least 3 pages (Cover + TOC + Chapters), got {len(r.pages)}"

    print("[PASS] Bug Y Test Passed: Comprehensive course bundling (Cover, TOC, multi-chapter PDF/DOCX/MD/Anki) verified with institution 'دانشگاه بق'.")

def test_bug_z_polymorphic_modules_and_balanced_coverage():
    """Bug Z: Polymorphic Module Architecture (Foundational Concept, Clinical Approach, Illness Script) and Anki Parity."""
    from build_summary_pack import build_docx_summary, build_markdown_summary, build_html_summary, get_or_synthesize_flashcards
    from audit_summary_coverage import audit_coverage, audit_illness_script_integrity

    data = {
        "topic_title": "فیزیوپاتولوژی جامع هیپوفیز و غدد",
        "topic_title_en": "Endocrine & Pituitary Comprehensive Review",
        "organ_system": "غدد درون‌ریز",
        "institution": "دانشگاه بق",
        "modules": [
            {
                "module_type": "foundational_concept",
                "name": "محور هیپوتالاموس-هیپوفیز-آدرنال (HPA Axis)",
                "anatomical_physiological_basis": [
                    "ترشح CRH از هسته پاراونتریکولار هیپوتالاموس و تحریک گیرنده‌های هیپوفیز قدامی"
                ],
                "regulatory_mechanisms_and_feedback": [
                    "مهار فیدبک منفی گلوکوکورتیکوئیدها (کورتیزول) بر ترشح CRH و ACTH در سطح مغز"
                ],
                "clinical_relevance": [
                    "قطع ناگهانی درمان طولانی‌مدت گلوکوکورتیکوئید منجر به نارسایی حاد ثانویه آدرنال می‌شود."
                ]
            },
            {
                "module_type": "clinical_approach",
                "name": "رویکرد تشخیصی گام‌به‌گام به سندرم کوشینگ",
                "chief_complaint_or_scenario": "چاقی مرکزی، استریای ارغوانی و هایپرتنشن مقاوم به درمان",
                "step_by_step_algorithm": [
                    {"step": "گام ۱: رد مصرف گلوکوکورتیکوئید اگزوژن", "action": "شرح حال دارویی دقیق و مکمل‌ها"},
                    {"step": "گام ۲: غربالگری افزایش کورتیزول", "action": "تست دگزامتازون ۱ میلی‌گرم شبانه یا کورتیزول آزاد ادرار ۲۴ ساعته"},
                    {"step": "گام ۳: سنجش ACTH پلاسما", "action": "افتراق وابستگی به ACTH از علل آدرنال خودمختار"}
                ],
                "high_yield_decision_points": [
                    "در صورت ACTH کمتر از 5 pg/mL علت آدرنال اولیه مطرح است و باید CT آدرنال انجام شود."
                ]
            },
            {
                "module_type": "illness_script",
                "name": "فئوکروموسیتوم (Pheochromocytoma)",
                "risk_factors_and_predisposing": ["سندرم MEN2، بیماری فون هیپل-لیندو (VHL) و NF-1"],
                "definition_and_etiology": ["تومور نورواندوکرین مترشحه کاتکول‌آمین‌ها از سلول‌های کرومافین مدولای آدرنال"],
                "pathophysiology_cascade": ["ترشح اپی‌نفرین و نوراپی‌نفرین -> تحریک گیرنده‌های آلفا و بتا -> وازوکانستریکشن شدید"],
                "clinical_presentation": [{"sign": "تریاد کلاسیک: سردرد، تعریق شدید و تپش قلب", "mechanism": "طوفان آدرنرژیک"}],
                "diagnostics": {
                    "best_initial_test": "سنجش متانفرین‌های آزاد پلاسما یا متانفرین‌های ادرار ۲۴ ساعته",
                    "gold_standard": "تصویربرداری تومور با MRI شکم یا اسکن MIBG"
                },
                "surgical_and_interventional_indications": [
                    "جراحی آدرنالکتومی با پیش‌درمان حداقل ۱۰ تا ۱۴ روزه با آلفا بلوکر (فنوکسی‌بنزامین)"
                ],
                "complications_and_prognosis": [
                    "بحران فشار خون و آریتمی‌های کشنده در صورت جراحی بدون بلوک آلفا و بتا"
                ],
                "contraindications_and_warnings": [
                    "بتابلوکر هرگز نباید قبل از بلوک کامل آلفا تجویز شود زیرا باعث بحران هایپرتنسیو مرگبار می‌شود."
                ],
                "pharmacology_and_management": [
                    {"drug_class": "فنوکسی‌بنزامین", "mechanism": "آلفا بلوکر غیرانتخابی غیرقابل برگشت", "exam_pearl": "شروع قبل از بتابلوکر"}
                ]
            }
        ]
    }

    out_docx = "dist/diag_test_z.docx"
    out_md = "dist/diag_test_z.md"
    out_html = "dist/diag_test_z.html"

    build_docx_summary(data, out_docx)
    build_markdown_summary(data, out_md)
    build_html_summary(data, out_html)

    # 1. Verify Markdown has polymorphic sections
    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()

    assert "مفاهیم و فیزیولوژی پایه" in md_text
    assert "رویکرد و الگوریتم تشخیصی" in md_text
    assert "زنجیره پاتوفیزیولوژی" in md_text
    assert "اندیکاسیون جراحی / مداخله‌ای" in md_text
    assert "عوارض وخیم و پیش‌آگهی" in md_text

    # 2. Verify Anki card auto-synthesis for all polymorphic branches
    cards = get_or_synthesize_flashcards(data)
    clozes = [c.get("text", "") or c.get("cloze_text", "") for c in cards]
    joined_clozes = " ".join(clozes)

    assert any("HPA" in c or "مهار فیدبک" in c or "CRH" in c for c in clozes), f"Foundational feedback Anki card missing! Cards: {clozes}"
    assert any("کوشینگ" in c or "ACTH" in c or "الگوریتم" in c for c in clozes), f"Clinical approach decision point Anki card missing! Cards: {clozes}"
    assert any("جراحی" in c or "فنوکسی‌بنزامین" in c or "آدرنالکتومی" in c for c in clozes), f"Surgical indication Anki card missing! Cards: {clozes}"
    assert any("عوارض" in c or "بحران" in c for c in clozes), f"Complication Anki card missing! Cards: {clozes}"

    # 3. Verify audit_illness_script_integrity evaluates all three types as sound modules
    audit_rep = audit_illness_script_integrity("", data)
    assert audit_rep["illness_script_compliant"], "Polymorphic modules failed illness script compliance!"
    assert audit_rep["sound_clinical_modules_count"] == 3
    assert not audit_rep["pharmacology_skew_detected"]

    print("[PASS] Bug Z Test Passed: Polymorphic module architecture (Foundational, Algorithm, Illness Script) and Anki parity verified.")

def test_bug_aa_high_compression_ratio_and_anti_bloat_audit():
    """Bug AA: High-Yield Compression Gate & Verbatim Bloat Prevention Audit."""
    from audit_summary_coverage import audit_coverage

    # Substantial clinical lecture text (>500 words)
    long_lecture_source = (
        "سندرم کوشینگ مجموعه‌ای از علائم بالینی ناشی از افزایش مزمن گلوکوکورتیکوئیدها در خون است. "
        "اتیولوژی این سندرم به دو گروه وابسته به ACTH و مستقل از ACTH تقسیم می‌شود. شایع‌ترین علت اگزوژن مصرف کورتون است. "
        "شایع‌ترین علت اندوژن، بیماری کوشینگ ناشی از میکروآدنوم هیپوفیز مترشحه ACTH است که حدود ۷۰ درصد موارد اندوژن را تشکیل می‌دهد. "
        "تومورهای آدرنال (آدنوم و کارسینوما) و ترشح نابجای ACTH توسط تومورهای نورواندوکرین نظیر کارسینوئید ریه از دیگر علل هستند. "
        "تظاهرات بالینی سندرم شامل چاقی مرکزی با تجمع چربی در صورت (چهره ماه)، بوفالو هامپ در پشت گردن، آتروفی عضلات پروگزیمال، "
        "پوست نازک و شکننده، استریای ارغوانی پهن با عرض بیش از ۱ سانتی‌متر بر روی شکم، اکیموز خودبه‌خودی، هایپرتنشن، "
        "اختلال تحمل گلوکز، استئوپروز شدید و افسردگی است. "
        "برای غربالگری و تشخیص افزایش کورتیزول سه تست اولیه وجود دارد: ۱) تست سرکوب شبانه دگزامتازون با دوز ۱ میلی‌گرم در ساعت ۲۳ "
        "و سنجش کورتیزول ساعت ۸ صبح که اگر کمتر از ۱.۸ میکروگرم در دسی‌لیتر باشد نرمال است؛ ۲) سنجش کورتیزول آزاد ادرار ۲۴ ساعته "
        "که اگر ۳ برابر حد نرمال باشد مثبت است؛ ۳) سنجش کورتیزول بزاق در نیمه‌شب که افزایش آن نشانه اختلال ریتم شبانه‌روزی است. "
        "پس از اثبات هایپرکورتیزولیسم، سطح ACTH پلاسما اندازه‌گیری می‌شود. اگر ACTH زیر ۵ پیکوگرم در میلی‌لیتر باشد مستقل از ACTH است "
        "و نیازمند سی‌تی‌اسکن غده آدرنال می‌باشد. اگر ACTH بالای ۲۰ پیکوگرم در میلی‌لیتر باشد وابسته به ACTH بوده و گام بعدی تست سرکوب "
        "با دوز بالای دگزامتازون (۸ میلی‌گرم) یا تست تحریک با CRH است. در بیماری کوشینگ هیپوفیزی، سطح کورتیزول بیش از ۵۰ درصد سرکوب می‌شود "
        "در حالی که در ترشح نابجای ACTH سرکوب دیده نمی‌شود. گلد استاندارد برای تایید آدنوم هیپوفیز، کاتتریزاسیون سینوس پتروزال تحتانی (BIPSS) "
        "با گرادیان مرکزی به محیطی بیش از ۲ قبل از CRH و بیش از ۳ بعد از تحریک CRH است. "
        "درمان خط اول در بیماری کوشینگ، جراحی ترانس اسفنوئیدال (TSS) آدنوم هیپوفیز است. در تومورهای آدرنال، آدرنالکتومی لاپاروسکوپیک انجام می‌شود. "
        "درمان‌های دارویی مهارکننده سنتز کورتیکواستروئید شامل کتوکونازول با دوز ۲۰۰ تا ۴۰۰ میلی‌گرم، متیراپون و متوتان در موارد عود یا غیرقابل جراحی "
        "کاربرد دارند. هشدار بسیار مهم: قطع ناگهانی درمان کورتیزول یا جراحی بدون جایگزینی استروئید منجر به نارسایی حاد آدرنال و شوک کشنده می‌شود. "
        "پایش بعد از عمل جراحی شامل اندازه‌گیری کورتیزول صبحگاهی است؛ افت سطح کورتیزول به زیر ۲ میکروگرم نشانه بهبودی موفق و پاکسازی آدنوم است."
    )

    # 1. High-Yield Compressed Summary (~100-150 words): preserves cutoffs, pillars, warnings
    compressed_summary = {
        "modules": [
            {
                "module_type": "illness_script",
                "name": "سندرم کوشینگ (Cushing Syndrome)",
                "definition_and_etiology": ["افزایش مزمن گلوکوکورتیکوئیدها؛ شایع‌ترین اگزوژن کورتون و اندوژن بیماری کوشینگ (۷۰٪)"],
                "pathophysiology_cascade": ["افزایش کورتیزول -> کاتابولیسم پروتئین، گلوکونئوژنز و مهار استئوبلاست‌ها"],
                "clinical_presentation": [{"sign": "چاقی مرکزی، چهره ماه، بوفالو هامپ، استریای ارغوانی > ۱ سانتی‌متر", "mechanism": "کاتابولیسم بافتی"}],
                "diagnostics": {
                    "best_initial_test": "تست دگزامتازون ۱ میلی‌گرم شبانه (کات‌آف ۱.۸ میکروگرم) یا کورتیزول ادرار ۲۴ ساعته",
                    "gold_standard": "کاتتریزاسیون سینوس پتروزال تحتانی (BIPSS) با گرادیان > ۲ یا > ۳ بعد از CRH"
                },
                "high_yield_pearls": [
                    "کات‌آف ACTH: زیر ۵ نشانه آدرنال (سی‌تی‌اسکن) و بالای ۲۰ وابسته به ACTH.",
                    "تست دگزامتازون ۸ میلی‌گرم: سرکوب > ۵۰٪ نشانه بیماری کوشینگ هیپوفیزی است."
                ],
                "surgical_and_interventional_indications": ["خط اول درمان: جراحی ترانس اسفنوئیدال (TSS) هیپوفیز یا آدرنالکتومی"],
                "contraindications_and_warnings": [
                    "هشدار بسیار مهم: قطع ناگهانی کورتون یا جراحی بدون جایگزینی استروئید منجر به بحران نارسایی حاد آدرنال و شوک کشنده می‌شود."
                ],
                "pharmacology_and_management": [
                    {"drug_class": "کتوکونازول و متیراپون", "mechanism": "مهار سنتز کورتیکواستروئید با دوز ۲۰۰ تا ۴۰۰ میلی‌گرم", "exam_pearl": "در موارد عود یا غیرقابل جراحی"}
                ]
            }
        ]
    }

    rep_compressed = audit_coverage(long_lecture_source, compressed_summary)
    assert rep_compressed["verdict"] == "PASSED", f"High-yield summary unexpectedly failed audit: {rep_compressed}"
    assert rep_compressed["scores"]["compression_ratio"] < 0.85, "Expected compact compression ratio!"
    assert not rep_compressed["scores"]["verbatim_bloat_detected"], "Bloat falsely detected on compact summary!"
    assert rep_compressed["compression_audit"]["high_yield_compressed"], "High yield compression flag not set!"

    # 2. Verbatim Bloat Summary: dumps entire raw lecture paragraphs verbatim into bullets
    bloated_summary = {
        "modules": [
            {
                "module_type": "illness_script",
                "name": "سندرم کوشینگ",
                "definition_and_etiology": [long_lecture_source[:300]],
                "pathophysiology_cascade": [long_lecture_source[300:600]],
                "clinical_presentation": [{"sign": long_lecture_source[600:900], "mechanism": "نامشخص"}],
                "diagnostics": {
                    "best_initial_test": long_lecture_source[900:1200],
                    "gold_standard": long_lecture_source[1200:1500]
                },
                "high_yield_pearls": [long_lecture_source[1500:1800]],
                "surgical_and_interventional_indications": [long_lecture_source[1800:2000]],
                "contraindications_and_warnings": [
                    "هشدار بسیار مهم: قطع ناگهانی درمان کورتیزول یا جراحی بدون جایگزینی استروئید منجر به نارسایی حاد آدرنال و شوک کشنده می‌شود."
                ],
                "pharmacology_and_management": [
                    {"drug_class": "کتوکونازول", "mechanism": long_lecture_source[2000:], "exam_pearl": "دوز ۲۰۰ تا ۴۰۰ میلی‌گرم"}
                ]
            }
        ]
    }

    rep_bloat = audit_coverage(long_lecture_source, bloated_summary)
    assert rep_bloat["scores"]["verbatim_bloat_detected"], "Failed to detect verbatim bloat on huge raw dump!"
    assert rep_bloat["verdict"] == "WARNING_VERBATIM_BLOAT_DETECTED", f"Expected bloat warning, got {rep_bloat['verdict']}"

    print("[PASS] Bug AA Test Passed: High-yield compression gate & verbatim bloat prevention audit verified.")

def test_bug_ab_ensure_master_matrices_polymorphic_parity_and_checklist():
    """Bug AB: Polymorphic Module Parity in Crash Sheet & Checklist Auto-Synthesis."""
    from build_summary_pack import ensure_master_matrices

    raw_data = {
        "topic_title": "فیزیوپاتولوژی جامع",
        "modules": [
            {
                "type": "foundational_concept",
                "name": "محور هورمونی HPA",
                "feedback_cascade": ["ترشح CRH -> ACTH -> کورتیزول -> مهار بازخوردی"],
                "physiological_role": ["تنظیم هموستاز گلوکز و پاسخ به استرس"],
                "baseline_tests": ["کورتیزول پلاسما 8 صبح: 5-25 mcg/dL"],
                "clinical_relevance": ["قطع ناگهانی درمان کورتون منجر به نارسایی حاد آدرنال می‌شود"]
            },
            {
                "type": "clinical_approach",
                "name": "رویکرد به شوک آنافیلاکسی",
                "initial_presentation": "کهیر، آنژیوادم، برونکواسپاسم و هایپوتنشن شدید",
                "step_by_step_algorithm": ["گام ۱: تزریق فوری اپی‌نفرین عضلانی ۰.۳ تا ۰.۵ میلی‌گرم", "گام ۲: اکسیژن‌تراپی و سرم نرمال‌سالین"],
                "pivot_points": ["در صورت عدم پاسخ بعد از ۵ دقیقه: تکرار اپی‌نفرین IM"],
                "definitive_action": "پایش حداقل ۶ تا ۸ ساعت جهت واکنش دو فازی"
            }
        ]
    }

    ensure_master_matrices(raw_data)

    assert "last_minute_crash_sheet" in raw_data, "Crash sheet was not auto-synthesized!"
    crash_rows = raw_data["last_minute_crash_sheet"]["rows"]
    assert len(crash_rows) == 2, f"Expected 2 rows in crash sheet, got {len(crash_rows)}"

    # Verify foundational row is NOT filled with dashes
    found_row = crash_rows[0]
    assert found_row[0] == "محور هورمونی HPA"
    assert found_row[1] == "هموستاز فیزیولوژیک"
    assert "CRH" in found_row[2] or "کورتیزول" in found_row[2]
    assert "استرس" in found_row[3] or "گلوکز" in found_row[3]
    assert "5-25" in found_row[4]
    assert "نارسایی" in found_row[5] or "کورتون" in found_row[5]

    # Verify approach row is NOT filled with dashes
    app_row = crash_rows[1]
    assert app_row[0] == "رویکرد به شوک آنافیلاکسی"
    assert "رویکرد" in app_row[1]
    assert "اپی‌نفرین" in app_row[2]
    assert "کهیر" in app_row[3]
    assert "تکرار" in app_row[4]
    assert "پایش" in app_row[5]

    # Verify morning checklist incorporated clinical relevance and pivot points
    assert "morning_exam_checklist" in raw_data, "Checklist was not auto-synthesized!"
    chk = raw_data["morning_exam_checklist"]
    assert any("کورتون" in b or "نارسایی" in b for b in chk), "Checklist missed foundational clinical relevance!"
    assert any("اپی‌نفرین" in b or "پایش" in b for b in chk), "Checklist missed approach pivot points!"

    print("[PASS] Bug AB Test Passed: Polymorphic module parity in crash sheet and checklist synthesis verified.")

def test_bug_ac_staging_classification_and_diagnostic_criteria_support():
    """Bug AC: Diagnostic Criteria & Staging Systems Support Across Formats."""
    from build_summary_pack import build_docx_summary, build_markdown_summary, build_html_summary, get_or_synthesize_flashcards

    data = {
        "topic_title": "کاردیولوژی بالینی",
        "modules": [
            {
                "type": "illness_script",
                "name": "اندوکاردیت عفونی (Infective Endocarditis)",
                "definition_and_etiology": ["عفونت دریچه‌های قلبی"],
                "staging_and_classification": {
                    "title": "معیارهای تشخیصی دوک (Duke Criteria)",
                    "headers": ["نوع معیار", "اجزا و نشانه‌ها", "تعریف قطعی"],
                    "rows": [
                        ["معیار ماژور", "کشت خون مثبت پایدار یا شواهد درگیری اندوکارد در اکو", "۲ ماژور یا ۱ ماژور + ۳ مینور"],
                        ["معیار مینور", "تب بالای ۳۸ درجه، مصرف داروی تزریقی، پدیده‌های عروقی و ایمونولوژیک", "۵ مینور"]
                    ]
                },
                "diagnostics": {"best_initial_test": "اکوکاردیوگرافی ترنس‌توراسیک (TTE)"}
            },
            {
                "type": "illness_script",
                "name": "آمبولی ریه (Pulmonary Embolism)",
                "definition_and_etiology": ["انسداد شریان ریوی با لخته"],
                "staging_and_classification": [
                    {"معیار": "علائم بالینی DVT", "امتیاز": "۳ امتیاز"},
                    {"معیار": "عدم وجود تشخیص محتمل‌تر", "امتیاز": "۳ امتیاز"},
                    {"معیار": "ضربان قلب بالای ۱۰۰", "امتیاز": "۱.۵ امتیاز"}
                ],
                "diagnostics": {"best_initial_test": "سی‌تی آنژیوگرافی ریه (CTPA)"}
            }
        ]
    }

    out_docx = "dist/diag_test_ac.docx"
    out_md = "dist/diag_test_ac.md"
    out_html = "dist/diag_test_ac.html"

    build_docx_summary(data, out_docx)
    build_markdown_summary(data, out_md)
    build_html_summary(data, out_html)

    # 1. Verify Markdown contains both criteria tables
    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()
    assert "معیارهای تشخیصی دوک (Duke Criteria)" in md_text
    assert "معیار ماژور" in md_text
    assert "علائم بالینی DVT" in md_text
    assert "۳ امتیاز" in md_text

    # 2. Verify HTML contains tables
    with open(out_html, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "Duke Criteria" in html_text
    assert "DVT" in html_text

    # 3. Verify Anki cards generated
    cards = get_or_synthesize_flashcards(data)
    clozes = [c.get("text", "") for c in cards]
    assert any("Duke" in c or "دوک" in c or "ماژور" in c for c in clozes), "Anki criteria card for Duke missing!"
    assert any("DVT" in c or "امتیاز" in c for c in clozes), "Anki criteria card for Wells missing!"

    print("[PASS] Bug AC Test Passed: Diagnostic criteria and staging systems rendered across DOCX, MD, HTML, and Anki.")

def test_bug_ad_emergency_and_special_populations_support():
    """Bug AD: Emergency Stabilization & Special Populations Support Across Formats."""
    from build_summary_pack import build_docx_summary, build_markdown_summary, build_html_summary, get_or_synthesize_flashcards

    data = {
        "topic_title": "اورژانس‌های داخلی و اندوکرین",
        "modules": [
            {
                "type": "illness_script",
                "name": "طوفان تیروئیدی (Thyroid Storm)",
                "emergency_management": [
                    "پایدارسازی فوری: تجویز پروپرانولول وریدی یا خوراکی جهت کنترل طوفان آدرنرژیک",
                    "تجویز هیدروکورتیزون وریدی جهت پیشگیری از نارسایی نسبی آدرنال و مهار تبدیل محیطی T4 به T3"
                ],
                "special_populations_and_pregnancy": [
                    "در سه ماهه اول بارداری: پروپیل‌تیواوراسیل (PTU) داروی انتخابی است.",
                    "از سه ماهه دوم و سوم به بعد: سوئیچ به متی‌مازول جهت کاهش سمیت کبدی شدید PTU."
                ],
                "contraindications_and_warnings": [
                    "محلول ید (Lugol) باید حداقل ۱ ساعت بعد از آنتی‌تیروئید تجویز شود تا اثر Jod-Basedow رخ ندهد."
                ]
            }
        ]
    }

    out_docx = "dist/diag_test_ad.docx"
    out_md = "dist/diag_test_ad.md"
    out_html = "dist/diag_test_ad.html"

    build_docx_summary(data, out_docx)
    build_markdown_summary(data, out_md)
    build_html_summary(data, out_html)

    # 1. Verify Markdown
    with open(out_md, "r", encoding="utf-8") as f:
        md_text = f.read()
    assert "اقدامات اورژانس و احیای فوری" in md_text
    assert "پروپرانولول" in md_text
    assert "ملاحظات گروه‌های خاص، بارداری و سالمندان" in md_text
    assert "سه ماهه اول بارداری" in md_text

    # 2. Verify HTML callouts
    with open(out_html, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "callout-emergency" in html_text
    assert "callout-special" in html_text
    assert "PTU" in html_text

    # 3. Verify Anki cards
    cards = get_or_synthesize_flashcards(data)
    clozes = [c.get("text", "") for c in cards]
    assert any("پروپرانولول" in c or "اورژانس" in c for c in clozes), "Emergency Anki card missing!"
    assert any("PTU" in c or "بارداری" in c for c in clozes), "Special population pregnancy Anki card missing!"

    print("[PASS] Bug AD Test Passed: Emergency management and special populations rendered across DOCX, MD, HTML, and Anki.")

def test_bug_ae_universal_dynamic_syndrome_and_disease_detection():
    """Bug AE: Universal Dynamic Syndrome and Disease Detection Across All Medical Specialties."""
    from audit_summary_coverage import audit_coverage, audit_illness_script_integrity

    source_text = (
        "در بخش قلب و عروق اطفال، بیماری کاوازاکی یکی از علل مهم واسکولیت است که منجر به آنوریسم عروق کرونر می‌شود. "
        "همچنین در شانت‌های چپ به راست درمان‌نشده، سندرم آیزن‌منگر در اثر وازوکانستریکشن ریوی ایجاد می‌گردد. "
        "در تامپوناد قلبی تریاد بک شامل هایپوتنشن، اتساع وریدهای ژوگولار و صداهای قلبی مبهم دیده می‌شود. "
        "برای تشخیص اندوکاردیت از معیارهای دوک و در آمبولی ریه از Wells Score استفاده می‌گردد. "
        "دوز اولیه آسپیرین ۸۰ میلی گرم بر کیلوگرم است."
    )

    # Incomplete summary that omitted Kawasaki disease and Eisenmenger syndrome
    incomplete_summary = {
        "modules": [
            {
                "type": "illness_script",
                "name": "تامپوناد قلبی",
                "clinical_presentation": [{"sign": "تریاد بک", "mechanism": "تجمع مایع پریکارد"}],
                "diagnostics": {"best_initial_test": "اکوکاردیوگرافی"}
            }
        ]
    }

    rep = audit_illness_script_integrity(source_text, incomplete_summary)
    omitted = rep.get("omitted_subtle_syndromes", [])
    detected = rep.get("detected_source_syndromes", [])

    # Verify that dynamic syndromes were detected from source
    assert any("کاوازاکی" in s for s in detected), f"Kawasaki disease was not detected from source! Detected: {detected}"
    assert any("آیزن" in s and "منگر" in s for s in detected), f"Eisenmenger syndrome was not detected from source! Detected: {detected}"
    assert any("Wells" in s or "wells" in s.lower() for s in detected), f"Wells Score was not detected from source! Detected: {detected}"

    # Verify that omitted syndromes were flagged
    assert any("کاوازاکی" in s for s in omitted), f"Kawasaki disease omission was not flagged! Omitted: {omitted}"
    assert any("آیزن" in s and "منگر" in s for s in omitted), f"Eisenmenger syndrome omission was not flagged! Omitted: {omitted}"
    assert not rep["illness_script_compliant"], "Audit should fail compliance when dynamic syndromes are omitted!"

    # Now verify complete summary passes
    complete_summary = {
        "modules": [
            {
                "type": "illness_script",
                "name": "بیماری کاوازاکی",
                "pathophysiology_cascade": ["واسکولیت عروق کرونر"],
                "clinical_presentation": [{"sign": "تب و بثورات", "mechanism": "واسکولیت"}],
                "diagnostics": {"best_initial_test": "اکوکاردیوگرافی"},
                "pharmacology_and_management": [{"drug_class": "آسپیرین", "mechanism": "ضد التهاب"}]
            },
            {
                "type": "illness_script",
                "name": "سندرم آیزن‌منگر و تامپوناد قلبی",
                "pathophysiology_cascade": ["معکوس شدن شانت در اثر افزایش فشار ریوی"],
                "clinical_presentation": [{"sign": "سیانوز و تریاد بک", "mechanism": "شانت راست به چپ"}],
                "diagnostics": {"best_initial_test": "معیارهای دوک و Wells Score"}
            }
        ]
    }
    rep_pass = audit_illness_script_integrity(source_text, complete_summary)
    assert rep_pass["illness_script_compliant"], f"Complete summary unexpectedly failed: {rep_pass['omitted_subtle_syndromes']}"
    assert len(rep_pass["omitted_subtle_syndromes"]) == 0, f"Expected 0 omissions, got {rep_pass['omitted_subtle_syndromes']}"

    print("[PASS] Bug AE Test Passed: Universal dynamic syndrome and disease detection across medical specialties verified.")

def test_bug_af_end_to_end_compact_high_yield_comprehensive_quad_pack():
    """Bug AF: Capstone end-to-end integration test verifying all 3 user requirements simultaneously:
    1. کم‌حجم (Ultra-condensed compression ratio 0.20 <= CR <= 0.65, zero verbatim bloat).
    2. High-Yield (Emergency protocols, lethal contraindications, diagnostic criteria/staging, special populations, Anki cloze cards).
    3. جامع (100% balanced coverage across physiology, clinical algorithms, and illness scripts with zero syndrome/disease omission).
    4. Quad-Pack Deliverables Parity (DOCX, MD, HTML, Playwright PDF, Anki TSV & APKG).
    """
    from build_summary_pack import (
        build_docx_summary,
        build_markdown_summary,
        build_html_summary,
        build_pdf_summary,
        build_anki_tsv,
        build_anki_apkg,
        ensure_master_matrices
    )
    from audit_summary_coverage import audit_coverage

    # Realistic 800+ word multi-topic lecture excerpt
    source_lecture = (
        "مقدمه و مروری بر پاتوفیزیولوژی سیستم غدد درون‌ریز و هورمون‌های تنظیم‌کننده آب و الکترولیت‌ها در بدن انسان: "
        "محور رنین-آنژیوتانسین-آلدوسترون یا همان سیستم RAAS یکی از حیاتی‌ترین و اصلی‌ترین مکانیسم‌های هموستاتیک بدن برای تنظیم طولانی‌مدت فشار خون شریانی، "
        "برون‌ده قلبی، حجم مؤثر در گردش خون و تعادل کاتیون‌های سدیم و پتاسیم و یون هیدروژن در نفرون‌های کلیوی است. در این مسیر بیوشیمیایی و سلولی پیچیده، "
        "هرگاه خون‌رسانی و پرفیوژن کلیوی به دلایلی نظیر هایپوولمی، دهیدراتاسیون، تنگی شریان کلیه یا افت فشار سیستمیک کاهش یابد، "
        "سلول‌های تخصص‌یافته ژوکستاگلومرولار (JG cells) در دیواره آرتریول آفرانت تحریک شده و آنزیم رنین را به داخل جریان خون آزاد می‌کنند. "
        "رنین پروتئین آنژیوتانسینوژن تولیدشده توسط کبد را به دکاپپتید غیرفعال آنژیوتانسین ۱ تبدیل می‌کند. سپس آنزیم مبدل آنژیوتانسین یا ACE عمدتاً "
        "در مویرگ‌های اندوتلیال عروق ریوی آن را به اکتاپپتید بسیار قدرتمند و فعال آنژیوتانسین ۲ تبدیل می‌نماید. آنژیوتانسین ۲ از طریق گیرنده‌های AT1 "
        "اثرات وسیعی اعمال می‌کند: اولاً موجب انقباض شدید عروق خونی محیطی و افزایش مقاومت عروقی سیستمیک می‌شود، ثانیاً قشر غده آدرنال و به‌ویژه منطقه "
        "گلومرولوزا را تحریک کرده تا هورمون استروئیدی آلدوسترون ترشح گردد. آلدوسترون بر روی سلول‌های اصلی (Principal cells) در لوله‌های جمع‌کننده قشری کلیه عمل کرده "
        "و با فعال‌سازی کانال‌های ENaC و پمپ سدیم-پتاسیم، بازجذب فعال سدیم و آب را افزایش داده و متقابلاً دفع پتاسیم و پروتون را به داخل ادرار تسریع می‌بخشد تا تعادل کاتیونی حفظ گردد. "
        "رویکرد بالینی و تشخیصی گام‌به‌گام به بیماران با هیپوناترمی: "
        "هیپوناترمی یکی از شایع‌ترین اختلالات الکترولیتی در بخش‌های بستری و اورژانس پزشکی است و به صورت سطح سدیم سرم کمتر از ۱۳۵ میلی‌اکی‌والان در لیتر تعریف می‌شود. "
        "گام اول و بنیادین در مواجهه با هر بیمار دچار هیپوناترمی، اندازه‌گیری دقیق اسمولالیته مؤثر سرم برای رد هیپوناترمی کاذب (مانند سودوهیپوناترمی ناشی از هایپرلیپیدمی شدید "
        "یا هایپرپروتئینمی و همچنین هیپوناترمی ترنس‌لوکاسیونال ناشی از هایپرگلیسمی شدید در دیابت کنترل‌نشده) است. در مرحله بعد پس از اثبات هیپوتونیک بودن، وضعیت حجم مایع "
        "خارج‌سلولی (ECF) بر اساس معاینه فیزیکی و سمیولوژی دقیق و تاریخچه بالینی به سه دسته کلی هیپوولمیک، یوولمیک و هایپرولمیک تقسیم‌بندی می‌شود. "
        "در هیپوناترمی هیپوولمیک سنجش سدیم ادرار کلید تمایز علل دفع کلیوی از خارج‌کلیوی است. در هیپوناترمی یوولمیک شایع‌ترین علت، سندرم ترشح نامتناسب هورمون آنتی‌دیورتیک یا SIADH است. "
        "پاتولوژی و تظاهرات بیماری‌های قشر آدرنال: "
        "سندرم کان یا آلدوسترونیسم اولیه (Primary Aldosteronism) در اثر ترشح خودمختار، نامتناسب و مهارنشدنی هورمون آلدوسترون توسط آدنوم یک‌طرفه قشر فوق کلیه یا هایپرپلازی "
        "دوطرفه کورتکس آدرنال ایجاد می‌شود. پاتوفیزیولوژی این اختلال شامل احتباس مزمن سدیم و اتساع حجم درون‌عروقی و هایپرتنشن شدید و اغلب مقاوم به درمان است، "
        "در حالی که دفع بیش از حد پتاسیم منجر به هایپوکالمی، ضعف عضلانی، خستگی و کرامپ و آلکالوز متابولیک ناشی از دفع پروتون می‌گردد. برای غربالگری تشخیصی استاندارد، "
        "سنجش نسبت غلظت آلدوسترون پلاسما به فعالیت رنین پلاسما یعنی PAC/PRA الزامی است و نسبت بالاتر از ۲۰ به همراه آلدوسترون بیش از ۱۵ نانوگرم بر دسی‌لیتر به عنوان نتیجه مثبت "
        "و مشکوک تلقی می‌شود. پس از مثبت شدن تست غربالگری، تست‌های تاییدی مهاری نظیر انفوزیون سالین ایزوتونیک یا تست بارگیری خوراکی سدیم برای اثبات خودمختار بودن ترشح انجام می‌گیرد. "
        "تصویربرداری انتخابی برای شناسایی ضایعه زمینه‌ای، سی‌تی اسکن آدرنال با مقاطع بسیار نازک است. برای درمان دارویی در هایپرپلازی دوطرفه یا قبل از جراحی، "
        "داروی اسپیرونولاکتون با دوز ۲۵ تا ۵۰ میلی‌گرم روزانه به عنوان آنتاگونیست رقابتی گیرنده مینرالوکورتیکوئید تجویز می‌شود و داروی اپلرنون جایگزین اختصاصی بدون عوارض ژنیکوماستی است. "
        "همچنین در تشخیص‌های افتراقی بیماری‌های ارثی توبول کلیوی، سندرم بارتر ناشی از موتاسیون در ناقل NKCC2 لوله هنله و سندرم گیتلمن ناشی از جهش در ترانسپورتر NCCT توبول دیستال باید مدنظر باشند. "
        "مدیریت اورژانس، کنترااندیکاسیون‌های خطرناک و احتیاطات بارداری: "
        "در بیماران با هیپوناترمی شدید و علامت‌دار حاد همراه با علائم تهدیدکننده حیات نظیر تشنج فعال، کاهش شدید سطح هوشیاری، کما یا ادم حاد مغزی، اقدام اورژانسی حیاتی "
        "تجویز فوری سالین ۳ درصد هیپرتونیک به میزان ۱۰۰ سی‌سی بولوس وریدی ظرف ۱۰ دقیقه است تا سریعاً از فتق مغزی جلوگیری شود. اما یک هشدار مرگبار و نکته امتحانی کلیدی که "
        "عدم رعایت آن به فاجعه منجر می‌شود این است که حداکثر سرعت اصلاح و تصحیح سدیم نباید از ۱۰ میلی‌اکی‌والان در ۲۴ ساعت فراتر برود؛ اصلاح بیش از حد سریع سدیم اکیداً ممنوع و کنترااندیکه قطعی است "
        "زیرا منجر به سندرم دمیلیناسیون اسموتیک مغزی (ODS / Central Pontine Myelinolysis) برگشت‌ناپذیر، کوما و آسیب دائمی نورولوژیک می‌شود. همچنین در دوران حاملگی، "
        "مصرف اسپیرونولاکتون به دلیل مهار گیرنده‌های آندروژن و خطر فمنیزاسیون جنین پسر مطلقاً کنترااندیکه بوده و باید داروی اپلرنون یا متیل‌دوپا به جای آن استفاده شود."
    )

    summary_data = {
        "topic_title": "فیزیوپاتولوژی سدیم و آلدوسترون",
        "topic_title_en": "Sodium and Aldosterone Pathophysiology",
        "organ_system": "نفرولوژی و غدد درون‌ریز",
        "modules": [
            {
                "type": "foundational_concept",
                "name": "محور رنین-آنژیوتانسین-آلدوسترون (RAAS)",
                "physiological_role": ["تنظیم فشار خون و هموستاز سدیم و پتاسیم"],
                "feedback_cascade": [
                    "کاهش پرفیوژن کلیه -> تحریک سلول‌های ژوکستاگلومرولار (JG cells) -> ترشح رنین",
                    "تبدیل آنژیوتانسینوژن به آنژیوتانسین ۱ توسط رنین و سپس آنژیوتانسین ۲ با آنزیم ACE",
                    "آنژیوتانسین ۲ از طریق گیرنده AT1 -> وازوکانستریکشن محیطی و ترشح آلدوسترون از قشر آدرنال",
                    "آلدوسترون با فعال‌سازی کانال ENaC -> بازجذب سدیم و آب و دفع ادراری پتاسیم و پروتون"
                ],
                "baseline_tests": ["سنجش رنین و آلدوسترون سرم"],
                "clinical_relevance": ["اساس پاتوفیزیولوژی هایپرتنشن رنوواسکولار و آلدوسترونیسم اولیه"],
                "high_yield_pearls": ["آنژیوتانسین ۲ قوی‌ترین وازوکانستریتور آندوژن سیستمیک است."]
            },
            {
                "type": "clinical_approach",
                "name": "رویکرد تشخیصی به هیپوناترمی",
                "initial_presentation": "سدیم سرم کمتر از ۱۳۵ میلی‌اکی‌والان",
                "step_by_step_algorithm": [
                    "گام ۱: اندازه‌گیری اسمولالیته سرم جهت رد سودوهیپوناترمی (هایپرلیپیدمی، هایپرپروتئینمی یا هایپرگلیسمی)",
                    "گام ۲: ارزیابی بالینی حجم مایع خارج‌سلولی (ECF): تفکیک به هیپوولمیک، یوولمیک (SIADH) و هایپرولمیک",
                    "گام ۳: سنجش سدیم ادرار برای افتراق دفع کلیوی از خارج‌کلیوی"
                ],
                "pivot_points": ["اسمولالیته سرم و سطح سدیم ادرار"],
                "definitive_action": "تصحیح اتیولوژی زمینه‌ای و بازگرداندن اسمولالیته نرمال",
                "emergency_management": {
                    "condition": "هیپوناترمی علامت‌دار حاد (تشنج، کما یا علائم شدید مغزی)",
                    "action": "تجویز فوری سالین ۳ درصد هیپرتونیک با دوز ۱۰۰ سی‌سی طی ۱۰ دقیقه",
                    "target": "مهار علائم حاد نورولوژیک و جلوگیری از فتق مغزی"
                },
                "contraindications_and_warnings": [
                    "هشدار مرگبار: سرعت اصلاح سدیم نباید از ۱۰ میلی‌اکی‌والان در ۲۴ ساعت بیشتر شود؛ اصلاح سریع به دلیل خطر سندرم دمیلیناسیون اسموتیک مغزی (ODS) و کوما اکیداً ممنوع است."
                ],
                "high_yield_pearls": ["اصلاح سدیم باید تدریجی و با پایش مداوم باشد."]
            },
            {
                "type": "illness_script",
                "name": "سندرم کان (آلدوسترونیسم اولیه)",
                "definition_and_etiology": ["ترشح خودمختار آلدوسترون ناشی از آدنوم یک‌طرفه یا هایپرپلازی قشر آدرنال"],
                "pathophysiology_cascade": [
                    "افزایش آلدوسترون -> احتباس سدیم و اتساع حجم -> هایپرتنشن شدید",
                    "دفع کلیوی پتاسیم و پروتون -> هایپوکالمی، ضعف عضلانی و آلکالوز متابولیک"
                ],
                "clinical_presentation": [
                    {"sign": "هایپرتنشن مقاوم، هایپوکالمی و کرامپ", "mechanism": "احتباس سدیم و اتلاف ادراری پتاسیم"}
                ],
                "staging_and_classification": [
                    {"stage": "غربالگری استاندارد", "criteria": "نسبت PAC/PRA بالاتر از ۲۰ همراه با آلدوسترون بیش از ۱۵ نانوگرم"},
                    {"stage": "تست تاییدی مهاری", "criteria": "تست بارگیری سدیم یا سالین ایزوتونیک جهت اثبات خودمختاری"},
                    {"stage": "تصویربرداری انتخابی", "criteria": "سی‌تی اسکن با کات نازک آدرنال"}
                ],
                "diagnostics": {
                    "best_initial_test": "نسبت آلدوسترون به رنین (PAC/PRA)",
                    "gold_standard": "تست تاییدی انفوزیون سالین و سی‌تی اسکن آدرنال"
                },
                "pharmacology_and_management": [
                    {
                        "drug_class": "اسپیرونولاکتون",
                        "mechanism": "آنتاگونیست رقابتی گیرنده مینرالوکورتیکوئید با دوز ۲۵ تا ۵۰ میلی‌گرم روزانه",
                        "indications": "هایپرپلازی دوطرفه کورتکس آدرنال یا پیش از جراحی",
                        "contraindications": "هایپرکالمی و نارسایی کلیه",
                        "exam_pearl": "درمان قطعی آدنوم یک‌طرفه جراحی لاپاروسکوپی است."
                    },
                    {
                        "drug_class": "اپلرنون",
                        "mechanism": "آنتاگونیست اختصاصی بدون اثرات فمنیزاسیون و ژنیکوماستی",
                        "indications": "عدم تحمل اسپیرونولاکتون یا بیماران باردار",
                        "contraindications": "هایپرکالمی",
                        "exam_pearl": "عدم بروز عوارض آنتی‌آندروژنیک نسبت به اسپیرونولاکتون."
                    }
                ],
                "special_populations_and_pregnancy": [
                    "مصرف اسپیرونولاکتون در بارداری به علت عوارض فمنیزاسیون جنین پسر اکیداً ممنوع بوده و داروی اپلرنون تجویز می‌گردد."
                ],
                "differential_diagnosis": [
                    {"disease": "سندرم بارتر", "distinguishing_feature": "موتاسیون در ترانسپورتر NKCC2 قوس لوله هنله"},
                    {"disease": "سندرم گیتلمن", "distinguishing_feature": "جهش در ترانسپورتر NCCT توبول دیستال"}
                ],
                "contraindications_and_warnings": [
                    "هشدار: مصرف اسپیرونولاکتون در حاملگی به دلیل مهار آندروژن جنین پسر کنترااندیکه قطعی است."
                ],
                "high_yield_pearls": [
                    "پدیده فرار آلدوسترون (Aldosterone Escape) مانع از بروز ادم محیطی بارز در سندرم کان می‌شود."
                ]
            }
        ]
    }

    # 1. Audit Coverage & Compression Validation
    audit_rep = audit_coverage(source_lecture, summary_data)
    assert audit_rep["verdict"] == "PASSED", f"Coverage audit failed: {audit_rep}"
    assert audit_rep["scores"]["numerical_coverage_pct"] >= 90.0, f"Low numerical coverage: {audit_rep['scores']['numerical_coverage_pct']}"
    assert audit_rep["scores"]["entity_recall_pct"] >= 90.0, f"Low entity recall: {audit_rep['scores']['entity_recall_pct']}"
    assert audit_rep["scores"]["contraindications_omitted_count"] == 0, "Lethal contraindications dropped!"
    assert audit_rep["scores"]["illness_script_compliant"], "Illness script integrity failed!"
    assert not audit_rep["scores"]["verbatim_bloat_detected"], "Bloat detected in ultra-condensed summary!"
    assert 0.20 <= audit_rep["scores"]["compression_ratio"] <= 0.65, f"Compression ratio out of range: {audit_rep['scores']['compression_ratio']}"

    # 2. Global Synthesis Validation (Master Matrices & Checklist)
    ensure_master_matrices(summary_data)
    assert "last_minute_crash_sheet" in summary_data, "Crash sheet missing!"
    assert len(summary_data["last_minute_crash_sheet"]["rows"]) == 3, "Crash sheet missing module rows!"
    assert "master_pharmacology_matrix" in summary_data, "Master pharm matrix missing!"
    assert len(summary_data["master_pharmacology_matrix"]["rows"]) == 2, "Pharm matrix missing drug rows!"
    assert "morning_exam_checklist" in summary_data, "Morning checklist missing!"
    assert len(summary_data["morning_exam_checklist"]) >= 4, "Morning checklist under-synthesized!"

    # 3. Deliverables Generation & Parity
    out_docx = "dist/diag_test_af_study_guide.docx"
    out_md = "dist/diag_test_af_summary.md"
    out_html = "dist/diag_test_af_view.html"
    out_pdf = "dist/diag_test_af_print.pdf"
    out_tsv = "dist/diag_test_af_anki.tsv"
    out_apkg = "dist/diag_test_af_anki.apkg"

    for p in [out_docx, out_md, out_html, out_pdf, out_tsv, out_apkg]:
        if os.path.exists(p):
            os.remove(p)

    build_docx_summary(summary_data, out_docx)
    build_markdown_summary(summary_data, out_md)
    build_html_summary(summary_data, out_html)
    build_pdf_summary(out_html, out_pdf, course_name="فیزیوپاتولوژی غدد و کلیه")
    build_anki_tsv(summary_data, out_tsv)
    build_anki_apkg(summary_data, out_apkg)

    # Deliverable file assertions
    assert os.path.exists(out_docx) and os.path.getsize(out_docx) > 5000, "DOCX generation failed!"
    assert os.path.exists(out_md) and os.path.getsize(out_md) > 1000, "Markdown generation failed!"
    assert os.path.exists(out_html) and os.path.getsize(out_html) > 1000, "HTML generation failed!"
    assert os.path.exists(out_pdf) and os.path.getsize(out_pdf) > 10000, "PDF generation failed!"
    assert os.path.exists(out_tsv) and os.path.getsize(out_tsv) > 500, "Anki TSV generation failed!"
    assert os.path.exists(out_apkg) and os.path.getsize(out_apkg) > 1000, "Anki APKG generation failed!"

    # Content assertions
    with open(out_md, "r", encoding="utf-8") as f_md:
        md_text = f_md.read()
    assert ("اقدامات اورژانس و احیای فوری" in md_text or "Emergency Resuscitation" in md_text), "Emergency protocol box missing in Markdown!"
    assert "سالین ۳ درصد" in md_text, "3% Saline missing in Markdown!"
    assert ("مرحله‌بندی" in md_text or "معیار" in md_text), "Diagnostic criteria missing in Markdown!"
    assert ("ملاحظات" in md_text and "بارداری" in md_text), "Pregnancy callout missing in Markdown!"
    assert ("کپسول طلایی شب امتحان" in md_text or "Cheat Sheet" in md_text), "Crash sheet missing in Markdown!"

    with open(out_html, "r", encoding="utf-8") as f_html:
        html_text = f_html.read()
    assert "دانشگاه بق" in html_text, "Institution 'دانشگاه بق' missing in HTML header!"
    assert "callout-emergency" in html_text, "Emergency callout CSS class missing in HTML!"

    with open(out_tsv, "r", encoding="utf-8") as f_tsv:
        tsv_lines = [l.strip() for l in f_tsv if l.strip()]
    assert len(tsv_lines) >= 15, f"Expected >= 15 Anki cloze cards, got {len(tsv_lines)}"
    assert all("{{c1::" in line for line in tsv_lines), "Not all Anki cards contain cloze deletions!"

    print("[PASS] Bug AF Test Passed: Capstone compact, high-yield, comprehensive quad-pack verified.")

def test_bug_ag_visual_diagram_and_multimodal_atlas_pipeline():
    """Bug AG: Visual Diagram Engine (High-DPI Mermaid rendering) and Multimodal Visual Asset Ingestion."""
    from build_summary_pack import (
        render_mermaid_diagram_to_png,
        build_docx_summary,
        build_markdown_summary,
        build_html_summary,
        get_or_synthesize_flashcards,
        build_anki_tsv,
        build_anki_apkg
    )
    from audit_summary_coverage import audit_coverage

    # 1. Test High-DPI Mermaid rendering to PNG via Playwright
    test_png = "dist/diag_test_ag_diagram.png"
    if os.path.exists(test_png):
        os.remove(test_png)

    mermaid_code = """
    graph TD
        CRH["هیپوتالاموس: ترشح CRH"] --> ACTH["هیپوفیز قدامی: ترشح ACTH"]
        ACTH --> Cortisol["کورتکس آدرنال: ترشح کورتیزول"]
        Cortisol -.->|"فیدبک منفی"| CRH
        Cortisol -.->|"فیدبک منفی"| ACTH
    """
    rendered = render_mermaid_diagram_to_png(mermaid_code, test_png, title="محور هیپوتالاموس-هیپوفیز-آدرنال")
    assert rendered, "Failed to render Mermaid diagram to PNG via Playwright!"
    assert os.path.exists(test_png) and os.path.getsize(test_png) > 1000, "Rendered PNG file missing or empty!"

    # 2. Test Multi-Format Deliverables with Diagrams and Visual Atlas
    data = {
        "topic_title": "نئوپلاسم‌های هیپوفیز و غدد درون‌ریز",
        "organ_system": "غدد درون‌ریز",
        "modules": [
            {
                "name": "آدنوم هیپوفیز (Pituitary Adenoma)",
                "diagram_image_path": test_png,
                "diagram_title": "الگوریتم رویکرد به توده هیپوفیز",
                "visual_atlas": [
                    {
                        "title": "گرافی MRI ساجیتال توده ماکروآدنوم هیپوفیز",
                        "image_path": test_png,
                        "clinical_findings": "توده سوپراسلار با قطر بالای ۱۰ میلی‌متر و فشار بر کیاسمای بینایی",
                        "exam_pearl": "شایع‌ترین تظاهر فشاری، نقص میدان بینایی دوطرفه تمپورال (Bitemporal Hemianopsia) است."
                    }
                ],
                "diagnostics": {
                    "best_initial_test": "سنجش سطح سرمی پرولاکتین و MRI هیپوفیز با گادولینیوم",
                    "gold_standard": "بررسی پاتولوژی پس از رزکسیون ترانس‌اسفنوئیدال (TSS)"
                },
                "high_yield_pearls": [
                    "میکروآدنوم (<۱۰mm) در برابر ماکروآدنوم (>=۱۰mm) نیازمند ارزیابی چشمی و پریمتری است."
                ]
            }
        ],
        "visual_atlas": [
            {
                "title": "علامت پمبرتون (Pemberton Sign) در گواتر غوطه‌ور",
                "image_path": test_png,
                "clinical_findings": "برافروختگی و پرخونی صورت (Facial Plethora) پس از بالا بردن بازوها به مدت ۱ دقیقه",
                "exam_pearl": "نشان‌دهنده انسداد توراسیک اینلت (Thoracic Inlet Obstruction) توسط گواتر رترواسترنال است."
            }
        ]
    }

    # 3. Test DOCX generation with diagrams and atlas
    out_docx = "dist/diag_test_ag.docx"
    build_docx_summary(data, out_docx)
    assert os.path.exists(out_docx) and os.path.getsize(out_docx) > 5000, "DOCX study guide with diagrams and atlas failed!"

    # 4. Test Markdown generation with diagrams and atlas
    out_md = "dist/diag_test_ag.md"
    build_markdown_summary(data, out_md)
    with open(out_md, "r", encoding="utf-8") as f_md:
        md_text = f_md.read()
    assert "📸 اطلس تصویری و تحلیل اسلایدهای بالینی" in md_text, "Visual Atlas missing in Markdown!"
    assert "گرافی MRI ساجیتال" in md_text, "MRI figure caption missing in Markdown!"
    assert "علامت پمبرتون" in md_text, "Top-level visual atlas missing in Markdown!"

    # 5. Test HTML generation with diagram styles and cards
    out_html = "dist/diag_test_ag.html"
    build_html_summary(data, out_html)
    with open(out_html, "r", encoding="utf-8") as f_html:
        html_text = f_html.read()
    assert "diagram-card" in html_text, "diagram-card class missing in HTML!"
    assert "figure-card" in html_text, "figure-card class missing in HTML!"
    assert "mermaid" in html_text, "Mermaid script/support missing in HTML!"
    assert "علامت پمبرتون" in html_text, "Pemberton figure missing in HTML!"

    # 6. Test Anki Flashcard Synthesis from Visual Atlas
    cards = get_or_synthesize_flashcards(data)
    atlas_cards = [c for c in cards if "VisualAtlas" in c.get("tag", "")]
    assert len(atlas_cards) >= 2, f"Expected at least 2 Visual Atlas Anki cards, got: {len(atlas_cards)}"
    assert any("MRI" in c["text"] or "ماکروآدنوم" in c["text"] for c in atlas_cards), "MRI atlas card missing!"
    assert any("پمبرتون" in c["text"] for c in atlas_cards), "Pemberton atlas card missing!"

    # 7. Test Audit Reporting of Visual Atlas
    src_text = "ماکروآدنوم هیپوفیز بالای ۱۰ میلی‌متر با درگیری کیاسمای بینایی. علامت پمبرتون در گواتر رترواسترنال."
    audit_rep = audit_coverage(src_text, data)
    assert "visual_atlas_audit" in audit_rep, "visual_atlas_audit missing from coverage report!"
    assert audit_rep["visual_atlas_audit"]["total_figures"] >= 2, "Visual atlas audit failed to count figures!"
    assert audit_rep["visual_atlas_audit"]["annotated_with_findings_or_pearls"] >= 2, "Visual atlas annotations not audited!"

    print("[PASS] Bug AG Test Passed: Visual diagram rendering and multimodal visual atlas pipeline verified across DOCX, MD, HTML, Anki, and Audit.")

def test_bug_ah_latex_bidi_markdown_and_flowchart_legibility_gates():
    """Bug AH: LaTeX/BiDi Formula Sanitization, True Word Bold Runs, Flowchart Legibility Gate & Quality Gates."""
    import docx
    from build_summary_pack import (
        sanitize_medical_text_symbols,
        format_inline_markdown,
        render_formula_html,
        clean_mermaid_code,
        render_mermaid_diagram_to_png,
        add_bidi_text,
        add_formula_card,
        add_markdown_alert,
        build_docx_summary,
        build_markdown_summary,
        build_html_summary
    )
    from audit_summary_coverage import (
        audit_raw_latex_leak,
        audit_raw_markdown_in_callouts,
        audit_markdown_alert_continuity
    )

    # 1. Test format_inline_markdown & sanitize_medical_text_symbols
    raw_sample = (
        r"۶. **دنوزوماب (Denosumab 120 mg SC):** آنتی‌بادی $\rightarrow$ افت کلسیم یونیزه و $\ge$ ۴.۰ mg/dL "
        r"با رسوب کریستال‌های $\text{Ca}_{10}(\text{PO}_4)_6(\text{OH})_2$ و هورمون $E_2$"
    )
    formatted_html = format_inline_markdown(raw_sample)
    assert "<strong>دنوزوماب (Denosumab 120 mg SC):</strong>" in formatted_html, "Markdown bold not converted to <strong>!"
    assert "**" not in formatted_html, f"Raw asterisks leaked in HTML: {formatted_html}"
    assert "←" in formatted_html, "TeX arrow not converted to Unicode!"
    assert "≥" in formatted_html, "TeX inequality not converted to Unicode!"
    assert "Ca₁₀(PO₄)₆(OH)₂" in formatted_html, "Chemical formula not converted to Unicode subscripts!"
    assert r"\text{" not in formatted_html, "Raw \\text{ leaked in HTML!"
    assert "E₂" in formatted_html, "TeX subscript not converted to Unicode subscript!"

    # 2. Test render_formula_html
    cccr_card = render_formula_html("CCCR = (U_Ca * S_Cr) / (S_Ca * U_Cr)", "فرمول کلیرانس کلسیم")
    assert "formula-card" in cccr_card, "formula-card class missing!"
    assert "formula-fraction" in cccr_card, "formula-fraction missing from CCCR!"
    assert "U<sub>Ca</sub>" in cccr_card, "U_Ca subscript missing!"

    ca_card = render_formula_html(r"\text{Corrected Calcium (mg/dL)} = \text{Measured Calcium} + 0.8 \times (4.0 - \text{Serum Albumin})")
    assert "formula-card" in ca_card, "formula-card missing from Corrected Ca!"
    assert r"\text{" not in ca_card, "Raw \\text{ leaked in formula card!"
    assert r"\times" not in ca_card, "Raw \\times leaked in formula card!"

    # 3. Test True Bold Run Creation in Word (.docx)
    doc = docx.Document()
    p = doc.add_paragraph()
    add_bidi_text(p, "۶. **دنوزوماب (Denosumab 120 mg SC):** آنتی‌بادی منوکلونال")
    
    # Assert runs exist, bold is preserved, and zero asterisks leak into Word runs
    bold_runs = [r for r in p.runs if r.bold and "دنوزوماب" in r.text]
    assert len(bold_runs) > 0, "True bold run not generated for **bold** text in Word!"
    for r in p.runs:
        assert "**" not in r.text, f"Raw '**' leaked in Word run text: '{r.text}'"

    # Test add_formula_card in Word
    add_formula_card(doc, "CCCR = (U_Ca * S_Cr) / (S_Ca * U_Cr)", "فرمول کلیرانس کلسیم")
    out_docx_ah = "dist/diag_test_ah.docx"
    doc.save(out_docx_ah)
    assert os.path.exists(out_docx_ah), "DOCX with formula card failed to save!"

    # 4. Test clean_mermaid_code & Flowchart Legibility Gate
    horizontal_code = "graph LR\n    A[\"تغییر pH $\\rightarrow$ کلسیم $\\ge$ ۱.۰\"] --> B[\"افت کلسیم یونیزه\"]"
    cleaned = clean_mermaid_code(horizontal_code)
    assert cleaned.startswith("flowchart TD"), f"Expected flowchart TD, got: {cleaned.splitlines()[0]}"
    assert "→" in cleaned or "←" in cleaned, "TeX arrow not sanitized in Mermaid code!"
    assert "≥" in cleaned, "TeX inequality not sanitized in Mermaid code!"
    assert r"\rightarrow" not in cleaned, "Raw \\rightarrow leaked in Mermaid code!"

    test_png_ah = "dist/diag_test_ah_mermaid.png"
    if os.path.exists(test_png_ah):
        os.remove(test_png_ah)
    rendered = render_mermaid_diagram_to_png(cleaned, test_png_ah, title="چارت عمودی خوانا")
    assert rendered, "Mermaid rendering failed!"
    assert os.path.exists(test_png_ah) and os.path.getsize(test_png_ah) > 1000, "Mermaid PNG missing!"

    # 5. Test Markdown Alert Continuity
    md_lines = []
    test_alert_items = [
        "خط اول احیا: سالین ۳ درصد سریع\nخط دوم احیا: دیورتیک لوپ پس از اتساع حجم",
        "پایش الکترولیت‌ها هر ۲ ساعت"
    ]
    add_markdown_alert(md_lines, "CAUTION", "پروتکل احیای اورژانسی", test_alert_items, icon="🚨")
    md_content = "\n".join(md_lines)
    
    # Check that EVERY line inside the alert has the '> ' prefix
    broken_lines = audit_markdown_alert_continuity(md_content)
    assert len(broken_lines) == 0, f"Markdown alert continuity broken: {broken_lines}"

    # 6. Test Quality Gates
    # LaTeX Leak Gate
    clean_text = "افزایش کلسیم ≥ ۱۲ mg/dL و هیدروکسی‌آپاتیت Ca₁₀(PO₄)₆(OH)₂"
    dirty_text = r"افزایش کلسیم \ge 12 mg/dL و \text{Ca}_{10}(\text{PO}_4)_6(\text{OH})_2"
    assert len(audit_raw_latex_leak(clean_text)) == 0, "False positive in LaTeX leak audit!"
    assert len(audit_raw_latex_leak(dirty_text)) >= 2, "Failed to detect raw LaTeX tokens in dirty text!"

    # HTML Callout Asterisk Gate
    clean_html_callout = '<div class="callout callout-pearl"><div>⚡ <strong>دنوزوماب:</strong> داروی آنتی‌بادی</div></div>'
    dirty_html_callout = '<div class="callout callout-pearl"><div>⚡ **دنوزوماب:** داروی آنتی‌بادی</div></div>'
    assert len(audit_raw_markdown_in_callouts(clean_html_callout)) == 0, "False positive in callout asterisk audit!"
    assert len(audit_raw_markdown_in_callouts(dirty_html_callout)) == 1, "Failed to detect raw '**' in HTML callout!"

    print("[PASS] Bug AH Test Passed: LaTeX/BiDi sanitization, true Word bold runs, flowchart legibility, alert continuity and quality gates verified.")

if __name__ == "__main__":
    print("=== RUNNING FULL SKILL DIAGNOSTICS SUITE ===")
    test_bug_a_apkg_synthesis_when_empty()
    test_bug_b_persian_units_and_prefixes()
    test_bug_c_table_column_overflow()
    test_bug_d_pharmacology_keys_interoperability()
    test_bug_e_null_diagnostics_no_orphan_headers()
    test_bug_f_sdt_content_control_unrolling()
    test_bug_g_none_resilience_across_all_builders()
    test_bug_h_merge_batches_resilience()
    test_bug_i_anki_cloze_model_type_and_card_generation()
    test_bug_j_partition_empty_and_markdown_tables()
    test_bug_k_decimal_sentence_splitting_and_exact_number_matching()
    test_bug_l_empty_callout_box_suppression()
    test_bug_m_word_break_and_tab_preservation()
    test_bug_n_master_matrices_word_markdown_parity()
    test_bug_o_contraindication_false_negative_prevention()
    test_bug_p_thousands_and_arabic_decimal_commas()
    test_bug_q_anki_contraindication_and_alert_synthesis()
    test_bug_r_table_cell_textbox_preservation()
    test_bug_s_english_prose_stop_words_in_entity_recall()
    test_bug_t_html_and_pdf_generation_with_bidi_stamping()
    test_bug_u_anti_pharmacology_skew_illness_script_integrity()
    test_bug_v_clinical_semiology_track_and_anki_synthesis()
    test_bug_w_morning_exam_checklist_and_comparative_matrices()
    test_bug_x_pdf_designer_engine_dual_pass_and_geometry()
    test_bug_y_comprehensive_course_bundling()
    test_bug_z_polymorphic_modules_and_balanced_coverage()
    test_bug_aa_high_compression_ratio_and_anti_bloat_audit()
    test_bug_ab_ensure_master_matrices_polymorphic_parity_and_checklist()
    test_bug_ac_staging_classification_and_diagnostic_criteria_support()
    test_bug_ad_emergency_and_special_populations_support()
    test_bug_ae_universal_dynamic_syndrome_and_disease_detection()
    test_bug_af_end_to_end_compact_high_yield_comprehensive_quad_pack()
    test_bug_ag_visual_diagram_and_multimodal_atlas_pipeline()
    test_bug_ah_latex_bidi_markdown_and_flowchart_legibility_gates()
    print("=== ALL DIAGNOSTIC TESTS COMPLETED SUCCESSFULLY ===")
